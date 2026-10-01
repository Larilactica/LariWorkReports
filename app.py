import datetime
import hmac
import base64
import gzip
import io
import json
import numbers
import os
import random
import re
import sqlite3
import unicodedata
from contextlib import closing

import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Gestión de Naves Espaciales",
    page_icon="🛸",
    layout="wide",
)

# =====================================================================
# CONFIGURACIÓN
# =====================================================================
# Si en Secrets existen SUPABASE_URL y SUPABASE_KEY, todo se guarda en Supabase (recomendado en Streamlit Cloud).
# Si no, se usa este archivo local, que en Streamlit Cloud se borra al reiniciar la app.
DB_PATH = os.environ.get("YANEZ_DB", "reportes_yanez.db")

# Se muestra en la barra lateral para saber qué versión del código está corriendo en Streamlit Cloud.
APP_VERSION = "v6 · gráfico en el NS compromiso (pantalla, PDF y Excel)"

# Cómo se busca cada columna: primero por el nombre del encabezado (exacto), luego por palabras que
# contenga, y solo como último recurso por su posición (0 = columna A), validando que el contenido tenga sentido.
BTK_CAMPOS = {
    "RUTA": {"etiqueta": "Ruta", "exactos": ["Identificador ruta", "Id ruta", "Ruta"], "pos": 0, "req": False},
    "PATENTE": {"etiqueta": "Patente", "exactos": ["Identificador", "Patente", "Placa"], "pos": 1, "req": True},
    "ORDEN": {"etiqueta": "Orden / guía", "exactos": ["Orden", "Numero de orden", "Pedido", "Guia"], "pos": 2, "req": True},
    "CLIENTE": {"etiqueta": "Cliente", "exactos": ["Cliente"], "pos": 3, "req": True},
    "ESTADO": {"etiqueta": "Estado", "exactos": ["Estado"], "pos": 10, "req": True},
    "SUBESTADO": {"etiqueta": "Sub-estado", "exactos": ["Subestado", "Sub estado", "Sub-estado"], "pos": 11, "req": True},
    "USUARIO": {"etiqueta": "Usuario móvil", "exactos": ["Usuario movil"], "pos": 12, "req": False},
    "F_COMP": {"etiqueta": "Fecha compromiso", "exactos": ["Fechaentrega", "Fecha entrega", "Fecha compromiso",
                                                          "Fecha de compromiso"], "pos": 72, "req": False},
    "COMUNA": {"etiqueta": "Comuna", "exactos": ["CMN", "Comuna"], "pos": 74, "req": False},
    "CONDUCTOR": {"etiqueta": "Conductor (ruta A, B, C…)", "exactos": ["Conductor"], "pos": 71, "req": False,
                  "silencioso": True},
    # Solo los archivos mensuales traen la fecha de cada día (columna A «Fecha»)
    "FECHA": {"etiqueta": "Fecha del día", "exactos": ["Fecha", "Fecha ruta"], "pos": 0, "req": False,
              "silencioso": True},
}
PARIS_CAMPOS = {
    "ORDEN": {"etiqueta": "Orden / guía", "exactos": ["Orden", "Numero de orden", "Pedido", "Guia", "N pedido"],
              "pos": 2, "req": True},
    "ESTADO": {"etiqueta": "Estado", "exactos": ["Estado"], "pos": 6, "req": True},
    "SUBESTADO": {"etiqueta": "Sub-estado", "exactos": ["Subestado", "Sub estado", "Sub-estado"], "pos": 7,
                  "req": True},
}
HELA_CAMPOS = {
    "DRIVER": {"etiqueta": "Driver", "exactos": ["Driver", "Conductor"], "contiene": ["DRIVER", "CONDUCTOR"],
               "pos": 0, "req": False, "silencioso": True},
    "PATENTE": {"etiqueta": "Patente", "exactos": ["Patente", "Placa"], "contiene": ["PATENTE", "PLACA"],
                "pos": 1, "req": False, "silencioso": True},
    "NOMBRE_RUTA": {"etiqueta": "Nombre ruta", "exactos": ["Nombre ruta", "Ruta"], "contiene": ["NOMBRE RUTA", "RUTA"],
                    "pos": 2, "req": False},
    "COMUNAS": {"etiqueta": "Comunas", "exactos": ["Comunas", "Comuna"], "contiene": ["COMUNA"], "pos": 3, "req": True},
    "PUNTOS": {"etiqueta": "Total puntos", "exactos": ["Total puntos", "Puntos"], "contiene": ["PUNTOS"],
               "pos": 4, "req": False},
    "COMPROMISOS": {"etiqueta": "Compromisos", "exactos": ["Compromisos", "Compromiso"], "contiene": ["COMPROMISO"],
                    "pos": 5, "req": False},
    "FECHA": {"etiqueta": "Fecha del día", "exactos": ["Fecha", "Dia"], "pos": 0, "req": False, "silencioso": True},
}
ESTADOS_CONOCIDOS = {"ENTREGADO", "NO ENTREGADO", "RECOGIDO", "NO RECOGIDO"}
CAMPOS_VALIDADOS = {"PATENTE", "ORDEN", "CLIENTE", "ESTADO", "F_COMP", "CONDUCTOR", "FECHA"}

TARIFAS_BASE = {
    "LOS ANDES + PUTAENDO": 75400, "PETORCA": 71470, "PUTAENDO": 63920,
    "SAN ESTEBAN": 63340, "CALLE LARGA": 63220, "LOS ANDES": 63220,
    "SANTO DOMINGO": 62410, "CABILDO": 57640, "RINCONADA": 57520,
    "SANTA MARIA": 57410, "SAN FELIPE": 56940, "SAN ANTONIO": 56470,
    "CARTAGENA": 51410, "EL TABO": 51200, "EL QUISCO": 50900,
    "LA LIGUA": 45340, "PANQUEHUE": 44790, "ALGARROBO": 40840,
    "CATEMU": 38780, "LLAY LLAY": 38060, "LLAYLLAY": 38060, "LLAY-LLAY": 38060,
    "PAPUDO": 37890, "ZAPALLAR": 31670, "NOGALES": 31200,
    "CASABLANCA": 25990, "HIJUELAS": 25680, "LA CALERA": 25170,
    "LA CRUZ": 23710, "PUCHUNCAVI": 19250, "QUILLOTA": 18540,
    "COLLIGUAY": 17070, "QUINTERO": 13500, "VALPARAISO": 12180,
    "CONCON": 0, "CON CON": 0, "QUILPUE": 0, "VILLA ALEMANA": 0,
    "OLMUE": 0, "LIMACHE": 0, "VIÑA DEL MAR": 0,
}
COMBO = "LOS ANDES + PUTAENDO"

# Estados que cuentan como "gestionado" en el NS ácido y en todos los conteos de entregados.
# "RECOGIDO" (retiros París) cuenta como entregado; "NO RECOGIDO" cuenta como no entregado.
ESTADOS_GESTIONADOS = {"ENTREGADO", "RECOGIDO"}

# Reglas del cruce de calces: (texto en sub-estado Yáñez, texto en sub-estado París)
REGLAS_CALCE = [
    ("SIN MORADORES", "CLIENTE NO ESTA"),
    ("CLIENTE ANULA", "EXPECTATIVA"),
    ("DIRECCION NO ENCONTRADA", "DIRECCION ERRONEA"),
    ("DESPACHO ADELANTADO", "MOTIVOS CLIENTE"),
    ("REPROGRAMADO", "MOTIVOS CLIENTE"),
]


# =====================================================================
# UTILIDADES DE TEXTO Y FECHAS
# =====================================================================
def limpiar_texto(texto):
    """Mayúsculas, sin tildes, sin puntos ni comas y sin espacios repetidos."""
    if texto is None or pd.isna(texto):
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    t = t.upper().replace(".", "").replace(",", "")
    return re.sub(r"\s+", " ", t).strip()


def txt(serie):
    s = serie.astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    return s.replace({"nan": "", "None": "", "NaT": "", "<NA>": ""})


def clave_patente(p):
    return re.sub(r"[\s\-\.]", "", str(p).upper())


def fmt_fecha(x):
    if x is None or pd.isna(x):
        return ""
    return x.strftime("%d/%m/%Y")


def pct(a, b):
    return (a / b * 100) if b else 0.0


def fmt_pct(x):
    return f"{x:.2f}".replace(".", ",") + " %"


def hoy_chile():
    try:
        from zoneinfo import ZoneInfo
        return datetime.datetime.now(ZoneInfo("America/Santiago")).date()
    except Exception:
        return datetime.date.today()


# =====================================================================
# TARIFAS DE COMUNAS
# =====================================================================
TARIFAS = {limpiar_texto(k): v for k, v in TARIFAS_BASE.items()}
NOMBRES_PARCIAL = sorted([n for n in TARIFAS if n != COMBO], key=len, reverse=True)


def obtener_tarifa_comuna(comuna_str):
    """Devuelve (nombre, tarifa) o None si no se reconoce. Primero coincidencia exacta."""
    c = limpiar_texto(comuna_str)
    if not c:
        return None
    if c in TARIFAS:
        return c, TARIFAS[c]
    for n in NOMBRES_PARCIAL:
        if re.search(rf"\b{re.escape(n)}\b", c):
            return n, TARIFAS[n]
    return None


def comuna_mas_lejana(comunas_texto):
    """De una lista separada por comas devuelve (comuna, valor, lista_no_reconocidas)."""
    texto = "" if comunas_texto is None or pd.isna(comunas_texto) else str(comunas_texto)
    encontrados, desconocidas = [], []
    for item in re.split(r"\s+-\s+|[,;/]", texto):
        if not limpiar_texto(item):
            continue
        r = obtener_tarifa_comuna(item)
        if r:
            encontrados.append(r)
        else:
            desconocidas.append(item.strip())
    nombres = {n for n, _ in encontrados}
    if "LOS ANDES" in nombres and "PUTAENDO" in nombres:
        encontrados.append((COMBO, TARIFAS[COMBO]))
    if not encontrados:
        return "SIN COMUNA RECONOCIDA", 0, desconocidas
    nombre, valor = max(encontrados, key=lambda x: x[1])
    return nombre, valor, desconocidas


# =====================================================================
# ALMACENAMIENTO
# Supabase si hay SUPABASE_URL y SUPABASE_KEY en Secrets; si no, un archivo local SQLite.
# =====================================================================
TAM_PAGINA = 1000  # Supabase devuelve máximo 1.000 filas por consulta: se consulta por páginas


def leer_secreto(nombre, defecto=None):
    try:
        return st.secrets[nombre]
    except Exception:  # noqa: BLE001
        return os.environ.get(nombre, defecto)


def _fila_archivo(f):
    return {k: f.get(k) for k in ("tipo", "fecha", "nombre", "tam_kb", "contenido_b64", "subido")}


class AlmacenSQLite:
    nombre = "local"

    def __init__(self, ruta):
        self.ruta = ruta
        with closing(self._conn()):
            pass

    def _conn(self):
        con = sqlite3.connect(self.ruta)
        con.execute("CREATE TABLE IF NOT EXISTS archivos (tipo TEXT NOT NULL, fecha TEXT NOT NULL, nombre TEXT, "
                    "tam_kb REAL, contenido_b64 TEXT, subido TEXT, PRIMARY KEY (tipo, fecha))")
        con.execute("CREATE TABLE IF NOT EXISTS resultados (tipo TEXT NOT NULL, fecha TEXT NOT NULL, "
                    "payload TEXT NOT NULL, PRIMARY KEY (tipo, fecha))")
        con.commit()
        return con

    def guardar_archivo(self, fila):
        with closing(self._conn()) as con:
            con.execute("INSERT OR REPLACE INTO archivos (tipo, fecha, nombre, tam_kb, contenido_b64, subido) "
                        "VALUES (:tipo, :fecha, :nombre, :tam_kb, :contenido_b64, :subido)", _fila_archivo(fila))
            con.commit()

    def nombre_archivo(self, tipo, fecha):
        with closing(self._conn()) as con:
            r = con.execute("SELECT nombre FROM archivos WHERE tipo=? AND fecha=?", (tipo, fecha)).fetchone()
        return r[0] if r else None

    def leer_archivo(self, tipo, fecha):
        with closing(self._conn()) as con:
            r = con.execute("SELECT nombre, contenido_b64 FROM archivos WHERE tipo=? AND fecha=?",
                            (tipo, fecha)).fetchone()
        return {"nombre": r[0], "contenido_b64": r[1]} if r else None

    def guardar_resultado(self, tipo, fecha, payload):
        with closing(self._conn()) as con:
            con.execute("INSERT OR REPLACE INTO resultados (tipo, fecha, payload) VALUES (?,?,?)",
                        (tipo, fecha, json.dumps(payload, ensure_ascii=False)))
            con.commit()

    def cargar_resultados(self, tipo, f_ini, f_fin):
        with closing(self._conn()) as con:
            rows = con.execute("SELECT fecha, payload FROM resultados WHERE tipo=? AND fecha BETWEEN ? AND ? "
                               "ORDER BY fecha", (tipo, f_ini, f_fin)).fetchall()
        return [(f, json.loads(p)) for f, p in rows]

    def borrar_resultado(self, tipo, fecha):
        with closing(self._conn()) as con:
            con.execute("DELETE FROM resultados WHERE tipo=? AND fecha=?", (tipo, fecha))
            con.commit()

    def listar_archivos(self):
        with closing(self._conn()) as con:
            rows = con.execute("SELECT tipo, fecha, nombre, tam_kb, subido FROM archivos "
                               "ORDER BY fecha DESC, tipo").fetchall()
        return [dict(zip(("tipo", "fecha", "nombre", "tam_kb", "subido"), r)) for r in rows]

    def listar_fechas(self):
        with closing(self._conn()) as con:
            rows = con.execute("SELECT fecha FROM archivos UNION SELECT fecha FROM resultados "
                               "WHERE tipo != 'config' ORDER BY fecha DESC").fetchall()
        return [r[0] for r in rows]

    def borrar_fecha(self, fecha):
        with closing(self._conn()) as con:
            con.execute("DELETE FROM archivos WHERE fecha=?", (fecha,))
            con.execute("DELETE FROM resultados WHERE fecha=? AND tipo != 'config'", (fecha,))
            con.commit()

    def limpiar_antes_de(self, fecha):
        """Borra todo lo anterior a la fecha (no toca la configuración). Devuelve (archivos, reportes)."""
        with closing(self._conn()) as con:
            n_a = con.execute("SELECT COUNT(*) FROM archivos WHERE fecha < ?", (fecha,)).fetchone()[0]
            n_r = con.execute("SELECT COUNT(*) FROM resultados WHERE fecha < ? AND tipo != 'config'",
                              (fecha,)).fetchone()[0]
            con.execute("DELETE FROM archivos WHERE fecha < ?", (fecha,))
            con.execute("DELETE FROM resultados WHERE fecha < ? AND tipo != 'config'", (fecha,))
            con.commit()
        return n_a, n_r

    def volcar(self):
        with closing(self._conn()) as con:
            a = con.execute("SELECT tipo, fecha, nombre, tam_kb, contenido_b64, subido FROM archivos").fetchall()
            r = con.execute("SELECT tipo, fecha, payload FROM resultados").fetchall()
        return {
            "archivos": [dict(zip(("tipo", "fecha", "nombre", "tam_kb", "contenido_b64", "subido"), x)) for x in a],
            "resultados": [{"tipo": t, "fecha": f, "payload": json.loads(p)} for t, f, p in r],
        }

    def restaurar(self, dump):
        for a in dump["archivos"]:
            self.guardar_archivo(a)
        for r in dump["resultados"]:
            self.guardar_resultado(r["tipo"], r["fecha"], r["payload"])
        return len(dump["archivos"]), len(dump["resultados"])


class AlmacenSupabase:
    nombre = "supabase"

    def __init__(self, url=None, key=None, cliente=None):
        if cliente is None:
            from supabase import create_client
            cliente = create_client(url, key)
        self.db = cliente

    def _paginar(self, construir, tam=TAM_PAGINA):
        filas, desde = [], 0
        while True:
            datos = construir().range(desde, desde + tam - 1).execute().data or []
            filas.extend(datos)
            if len(datos) < tam:
                return filas
            desde += tam

    def guardar_archivo(self, fila):
        self.db.table("archivos").upsert(_fila_archivo(fila), on_conflict="tipo,fecha").execute()

    def nombre_archivo(self, tipo, fecha):
        d = (self.db.table("archivos").select("nombre").eq("tipo", tipo).eq("fecha", fecha)
             .limit(1).execute().data)
        return d[0]["nombre"] if d else None

    def leer_archivo(self, tipo, fecha):
        d = (self.db.table("archivos").select("nombre,contenido_b64").eq("tipo", tipo).eq("fecha", fecha)
             .limit(1).execute().data)
        return d[0] if d else None

    def guardar_resultado(self, tipo, fecha, payload):
        self.db.table("resultados").upsert({"tipo": tipo, "fecha": fecha, "payload": payload},
                                           on_conflict="tipo,fecha").execute()

    def cargar_resultados(self, tipo, f_ini, f_fin):
        filas = self._paginar(lambda: self.db.table("resultados").select("fecha,payload").eq("tipo", tipo)
                              .gte("fecha", f_ini).lte("fecha", f_fin).order("fecha"))
        out = []
        for r in filas:
            p = r["payload"]
            out.append((r["fecha"], json.loads(p) if isinstance(p, str) else p))
        return out

    def borrar_resultado(self, tipo, fecha):
        self.db.table("resultados").delete().eq("tipo", tipo).eq("fecha", fecha).execute()

    def listar_archivos(self):
        return self._paginar(lambda: self.db.table("archivos").select("tipo,fecha,nombre,tam_kb,subido")
                             .order("fecha", desc=True))

    def listar_fechas(self):
        fechas = set()
        for r in self._paginar(lambda: self.db.table("archivos").select("fecha").order("fecha")):
            fechas.add(r["fecha"])
        for r in self._paginar(lambda: self.db.table("resultados").select("fecha").neq("tipo", "config")
                               .order("fecha")):
            fechas.add(r["fecha"])
        return sorted(fechas, reverse=True)

    def borrar_fecha(self, fecha):
        self.db.table("archivos").delete().eq("fecha", fecha).execute()
        self.db.table("resultados").delete().eq("fecha", fecha).neq("tipo", "config").execute()

    def limpiar_antes_de(self, fecha):
        """Borra todo lo anterior a la fecha (no toca la configuración). Devuelve (archivos, reportes)."""
        n_a = len(self._paginar(lambda: self.db.table("archivos").select("tipo,fecha").lt("fecha", fecha)
                                .order("fecha")))
        n_r = len(self._paginar(lambda: self.db.table("resultados").select("tipo,fecha").lt("fecha", fecha)
                                .neq("tipo", "config").order("fecha")))
        self.db.table("archivos").delete().lt("fecha", fecha).execute()
        self.db.table("resultados").delete().lt("fecha", fecha).neq("tipo", "config").execute()
        return n_a, n_r

    def volcar(self):
        return {
            "archivos": self._paginar(lambda: self.db.table("archivos").select("*").order("fecha"), tam=20),
            "resultados": self._paginar(lambda: self.db.table("resultados").select("*").order("fecha"), tam=100),
        }

    def restaurar(self, dump):
        for i in range(0, len(dump["archivos"]), 20):
            lote = [_fila_archivo(a) for a in dump["archivos"][i:i + 20]]
            self.db.table("archivos").upsert(lote, on_conflict="tipo,fecha").execute()
        for i in range(0, len(dump["resultados"]), 50):
            lote = [{"tipo": r["tipo"], "fecha": r["fecha"], "payload": r["payload"]}
                    for r in dump["resultados"][i:i + 50]]
            self.db.table("resultados").upsert(lote, on_conflict="tipo,fecha").execute()
        return len(dump["archivos"]), len(dump["resultados"])


def crear_almacen():
    url, key = leer_secreto("SUPABASE_URL"), leer_secreto("SUPABASE_KEY")
    if url and key:
        return AlmacenSupabase(str(url), str(key))
    return AlmacenSQLite(DB_PATH)


STORE = None  # se asigna al iniciar la interfaz


def _invalidar():
    """La interfaz lo reemplaza para limpiar la caché cuando se guardan o borran datos."""


def _memo(clave, fn, *args):
    """La interfaz lo reemplaza por una versión con caché para no consultar la base en cada clic."""
    return fn(*args)


def _json_default(o):
    return o.item() if hasattr(o, "item") else str(o)


def guardar_archivo(tipo, fecha, nombre, contenido):
    STORE.guardar_archivo({
        "tipo": tipo, "fecha": fecha.isoformat(), "nombre": nombre,
        "tam_kb": round(len(contenido) / 1024, 1),
        "contenido_b64": base64.b64encode(contenido).decode("ascii"),
        "subido": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
    })
    _invalidar()


def archivo_guardado(tipo, fecha):
    r = STORE.leer_archivo(tipo, fecha.isoformat())
    return (r["nombre"], base64.b64decode(r["contenido_b64"])) if r else None


def nombre_archivo(tipo, fecha):
    return _memo("nombre_archivo", STORE.nombre_archivo, tipo, fecha.isoformat())


def resolver_archivo(tipo, fecha, subido):
    """Si se subió un archivo lo guarda con su fecha y lo usa; si no, usa el guardado para esa fecha."""
    if subido is not None:
        contenido = subido.getvalue()
        guardar_archivo(tipo, fecha, subido.name, contenido)
        return subido.name, contenido
    return archivo_guardado(tipo, fecha)


def guardar_resultado(tipo, fecha, meta=None, dfs=None):
    payload = {
        "meta": json.loads(json.dumps(meta or {}, default=_json_default)),
        "dfs": {k: json.loads(d.to_json(orient="split")) for k, d in (dfs or {}).items()},
    }
    STORE.guardar_resultado(tipo, fecha.isoformat(), payload)
    _invalidar()


ETIQUETAS_COMUNAS_ANTIGUAS = {'S: Coincide puntos': 'U: Coincide puntos', 'T: Compromisos BTK': 'V: Compromisos BTK', 'U: Coincide compromisos': 'W: Coincide compromisos', 'V: Usuario móvil BTK': 'X: Usuario móvil BTK', 'W: Comuna lejana según Hela': 'Y: Comuna lejana según Hela', 'X: Coincide con Hela': 'Z: Coincide con Hela', 'Y: Comunas no reconocidas': 'AA: Comunas no reconocidas', 'Z: Origen': 'AB: Origen', 'AA: Dif. puntos (BTK − Hela)': 'AC: Dif. puntos (BTK − Hela)', 'AB: Dif. compromisos (BTK − Hela)': 'AD: Dif. compromisos (BTK − Hela)', 'AC: Comunas agregadas desde BTK': 'AE: Comunas agregadas desde BTK', 'AD: Valor según Hela': 'AF: Valor según Hela', 'AE: Letra de ruta': 'AG: Letra de ruta', 'AF: Pedidos de otras rutas en esta patente': 'AH: Pedidos de otras rutas en esta patente', 'AG: Pedidos de esta ruta llevados por otra patente': 'AI: Pedidos de esta ruta llevados por otra patente', 'AH: Pedidos con la letra de esta ruta (BTK)': 'AJ: Pedidos con la letra de esta ruta (BTK)'}


def _migrar_comunas(d):
    """Pone al día el reporte de comunas guardado con el formato anterior (sin «Pedidos Easy / París»)."""
    if "S: Coincide puntos" in d.columns:
        d = d.rename(columns=ETIQUETAS_COMUNAS_ANTIGUAS)
        pos = list(d.columns).index("R: Puntos BTK") + 1
        d.insert(pos, "S: Pedidos Easy", d["M: Entregados Easy BTK"] + d["N: No Entregados Easy BTK"])
        d.insert(pos + 1, "T: Pedidos París", d["O: Entregados París BTK"] + d["P: No Entregados París BTK"])
    return d


def _enteros_donde_corresponde(d):
    """Las columnas numéricas cuyos valores son todos enteros (con huecos) se muestran sin decimales: 13, no 13.000000."""
    d = d.copy()
    for c in d.columns:
        if _es_pct(c) or not pd.api.types.is_numeric_dtype(d[c]) or pd.api.types.is_bool_dtype(d[c]):
            continue
        v = d[c].dropna()
        if len(v) and (v % 1 == 0).all():
            d[c] = d[c].astype("Int64")
    return d


def _cargar_sin_cache(tipo, f_ini, f_fin):
    out = []
    for f, payload in STORE.cargar_resultados(tipo, f_ini.isoformat(), f_fin.isoformat()):
        dfs = {k: pd.read_json(io.StringIO(json.dumps(v)), orient="split", dtype=False, convert_dates=False)
               for k, v in payload.get("dfs", {}).items()}
        for k in ("reporte", "comunas"):
            if k in dfs:
                dfs[k] = _enteros_donde_corresponde(_migrar_comunas(dfs[k]))
        out.append((datetime.date.fromisoformat(f), payload.get("meta", {}), dfs))
    return out


def cargar(tipo, f_ini, f_fin):
    """Devuelve lista de (fecha, meta, {nombre: DataFrame}) ordenada por fecha."""
    return _memo("cargar", _cargar_sin_cache, tipo, f_ini, f_fin)


def _listar_archivos_sin_cache():
    cols = ["Fecha", "Tipo", "Archivo", "Tamaño (KB)", "Subido el"]
    filas = [[f["fecha"], f["tipo"], f["nombre"], f["tam_kb"], f["subido"]] for f in STORE.listar_archivos()]
    return pd.DataFrame(filas, columns=cols)


def listar_archivos():
    return _memo("listar_archivos", _listar_archivos_sin_cache)


def listar_fechas():
    return _memo("listar_fechas", STORE.listar_fechas)


def borrar_fecha(fecha_iso):
    STORE.borrar_fecha(fecha_iso)
    _invalidar()


# --- Limpieza automática: el día D de cada mes se borra todo lo anterior al primer día de ese mes ---
CONFIG_FECHA = datetime.date(2999, 12, 31)  # fecha "reservada" donde se guarda la configuración
CONFIG_DEFECTO = {"activa": True, "dia": 5, "primera": "2026-11-05", "ultima": None}


def _leer_config_sin_cache():
    cfg = dict(CONFIG_DEFECTO)
    r = STORE.cargar_resultados("config", CONFIG_FECHA.isoformat(), CONFIG_FECHA.isoformat())
    if r:
        cfg.update(r[0][1].get("meta", {}))
    return cfg


def leer_config_limpieza(con_cache=True):
    return _memo("config_limpieza", _leer_config_sin_cache) if con_cache else _leer_config_sin_cache()


def guardar_config_limpieza(cfg):
    guardar_resultado("config", CONFIG_FECHA, {k: cfg.get(k) for k in CONFIG_DEFECTO}, {})


def limpieza_pendiente(cfg, hoy):
    """Fecha de corte (primer día del mes actual) si hoy corresponde limpiar; si no, None."""
    if not cfg.get("activa"):
        return None
    if hoy < datetime.date.fromisoformat(cfg["primera"]) or hoy.day < int(cfg["dia"]):
        return None
    ultima = datetime.date.fromisoformat(cfg["ultima"]) if cfg.get("ultima") else None
    if ultima and (ultima.year, ultima.month) == (hoy.year, hoy.month):
        return None  # este mes ya se limpió
    return datetime.date(hoy.year, hoy.month, 1)


def ejecutar_limpieza(hoy):
    """Borra lo anterior al mes en curso si corresponde. Devuelve (corte, archivos, reportes) o None."""
    cfg = leer_config_limpieza(con_cache=False)
    corte = limpieza_pendiente(cfg, hoy)
    if corte is None:
        return None
    n_a, n_r = STORE.limpiar_antes_de(corte.isoformat())
    cfg["ultima"] = hoy.isoformat()
    guardar_config_limpieza(cfg)
    _invalidar()
    return corte, n_a, n_r


def limpiar_antes_de(corte):
    """Limpieza manual: borra todo lo anterior a la fecha. Devuelve (archivos, reportes)."""
    n = STORE.limpiar_antes_de(corte.isoformat())
    _invalidar()
    return n


def proxima_limpieza(cfg, hoy):
    """(fecha de la próxima limpieza, fecha de corte) o None si está desactivada."""
    if not cfg.get("activa"):
        return None
    primera = datetime.date.fromisoformat(cfg["primera"])
    ultima = datetime.date.fromisoformat(cfg["ultima"]) if cfg.get("ultima") else None
    y, m = hoy.year, hoy.month
    for _ in range(36):
        f = datetime.date(y, m, int(cfg["dia"]))
        ya = ultima is not None and (ultima.year, ultima.month) == (y, m)
        if f >= primera and f >= hoy and not ya:
            return f, datetime.date(f.year, f.month, 1)
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return None


def exportar_respaldo():
    """Copia completa (archivos y reportes) en un .json.gz."""
    crudo = json.dumps(STORE.volcar(), ensure_ascii=False, default=_json_default).encode("utf-8")
    return gzip.compress(crudo)


def importar_respaldo(datos):
    try:
        dump = json.loads(gzip.decompress(datos).decode("utf-8"))
        ok = isinstance(dump, dict) and isinstance(dump.get("archivos"), list) \
            and isinstance(dump.get("resultados"), list)
    except Exception:  # noqa: BLE001
        ok = False
    if not ok:
        raise ValueError("El archivo no es un respaldo válido de esta app.")
    n = STORE.restaurar(dump)
    _invalidar()
    return n


# =====================================================================
# LECTURA Y PREPARACIÓN DE ARCHIVOS
# =====================================================================
def leer_archivo(nombre, contenido):
    if nombre.lower().endswith((".xlsx", ".xls")):
        return pd.read_excel(io.BytesIO(contenido))
    ultimo = None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(contenido), sep=None, engine="python", encoding=enc)
        except Exception as e:  # noqa: BLE001
            ultimo = e
    raise ultimo


def letra_excel(i):
    letras, n = "", i + 1
    while n:
        n, r = divmod(n - 1, 26)
        letras = chr(65 + r) + letras
    return letras


def _valida(campo, serie):
    """¿El contenido de la columna tiene sentido para este dato?"""
    if campo not in CAMPOS_VALIDADOS:
        return True
    v = txt(serie)
    v = v[v.ne("")]
    if v.empty:
        return False
    if campo == "PATENTE":
        return v.str.upper().str.match(r"^(?=.*[A-Z])(?=.*\d)[A-Z0-9\-\. ]{5,9}$").mean() >= 0.6
    if campo == "ORDEN":
        # en los archivos mensuales una orden puede repetirse en varios días, así que no se exige que sea única
        return v.str.match(r"^[A-Za-z0-9\-_/]{4,30}$").mean() >= 0.8
    if campo == "CLIENTE":
        return bool(v.map(limpiar_texto).str.contains("EASY|PARIS").any())
    if campo == "ESTADO":
        return bool(v.map(limpiar_texto).isin(ESTADOS_CONOCIDOS).any())
    if campo == "F_COMP":
        return pd.to_datetime(v, errors="coerce", format="mixed", dayfirst=True).notna().mean() >= 0.3
    if campo == "FECHA":
        if pd.api.types.is_numeric_dtype(serie):
            return False
        f = pd.to_datetime(v, errors="coerce", format="mixed", dayfirst=True)
        return bool(((f.dt.year >= 2000) & (f.dt.year <= 2100)).mean() >= 0.5)
    if campo == "CONDUCTOR":
        return v.str.contains(r"(?i)\bruta\s+[A-Z]{1,2}\b").mean() >= 0.3
    return True


def mapear_columnas(df, campos):
    """Encuentra cada columna: 1) encabezado exacto, 2) encabezado parcial, 3) posición (solo si el
    contenido calza). Las pasadas van en ese orden para que una posición no le robe la columna a otro dato."""
    cols = [limpiar_texto(c) for c in df.columns]
    usados, hallado = set(), {}
    for campo, spec in campos.items():
        for alias in spec["exactos"]:
            a = limpiar_texto(alias)
            idx = next((i for i, c in enumerate(cols) if c == a and i not in usados), None)
            if idx is not None:
                hallado[campo] = (idx, "encabezado")
                usados.add(idx)
                break
    for campo, spec in campos.items():
        if campo in hallado:
            continue
        for kw in spec.get("contiene", []):
            idx = next((i for i, c in enumerate(cols) if kw in c and i not in usados), None)
            if idx is not None:
                hallado[campo] = (idx, "encabezado (parcial)")
                usados.add(idx)
                break
    for campo, spec in campos.items():
        if campo in hallado:
            continue
        pos = spec["pos"]
        if pos is not None and pos < df.shape[1] and pos not in usados and _valida(campo, df.iloc[:, pos]):
            hallado[campo] = (pos, "posición (no se encontró el encabezado)")
            usados.add(pos)

    mapa, avisos, faltan = {}, [], []
    for campo, spec in campos.items():
        if campo not in hallado:
            if spec["req"]:
                faltan.append(spec["etiqueta"])
            elif not spec.get("silencioso"):
                avisos.append(f"No se encontró la columna de {spec['etiqueta']}; ese dato quedará vacío.")
            continue
        idx, metodo = hallado[campo]
        if (metodo != "posición (no se encontró el encabezado)" and not spec.get("silencioso")
                and not _valida(campo, df.iloc[:, idx])):
            avisos.append(f"La columna «{df.columns[idx]}» se usó como {spec['etiqueta']}, pero su contenido "
                          "no parece corresponder. Revísala.")
        mapa[campo] = {"etiqueta": spec["etiqueta"], "columna": str(df.columns[idx]),
                       "letra": letra_excel(idx), "indice": idx, "metodo": metodo}
    if faltan:
        detectados = ", ".join(str(c) for c in list(df.columns)[:25]) + (" …" if df.shape[1] > 25 else "")
        raise ValueError("No se encontraron estas columnas en el archivo: " + ", ".join(faltan)
                         + ". Encabezados del archivo: " + detectados)
    return mapa, avisos


def mapa_a_filas(mapa):
    return [{"Dato": m["etiqueta"], "Columna en el archivo": m["columna"], "Letra": m["letra"],
             "Cómo se encontró": m["metodo"]} for m in mapa.values()]


def _col(df, mapa, nombre):
    if nombre in mapa:
        return df.iloc[:, mapa[nombre]["indice"]]
    return pd.Series([None] * len(df), index=df.index, dtype="object")


def preparar_btk(df):
    """Ordena el BTK de Yáñez en columnas con nombre, buscándolas por encabezado."""
    mapa, avisos = mapear_columnas(df, BTK_CAMPOS)

    def col(nombre):
        return _col(df, mapa, nombre)

    b = pd.DataFrame(index=df.index)
    b["RUTA"] = txt(col("RUTA"))
    b["PATENTE"] = txt(col("PATENTE")).str.upper()
    b["PAT_KEY"] = b["PATENTE"].map(clave_patente)
    b["ORDEN"] = txt(col("ORDEN"))
    b["CLIENTE_TXT"] = txt(col("CLIENTE"))
    b["CLIENTE"] = b["CLIENTE_TXT"].map(limpiar_texto)
    b["ESTADO_TXT"] = txt(col("ESTADO"))
    b["ESTADO"] = b["ESTADO_TXT"].map(limpiar_texto)
    b["SUBESTADO"] = txt(col("SUBESTADO"))
    b["SUB_CLEAN"] = b["SUBESTADO"].map(limpiar_texto)
    b["USUARIO"] = txt(col("USUARIO"))
    b["F_COMP"] = pd.to_datetime(col("F_COMP"), errors="coerce", dayfirst=True, format="mixed").dt.date
    b["COMUNA"] = txt(col("COMUNA"))
    b["FECHA"] = pd.to_datetime(col("FECHA"), errors="coerce", dayfirst=True, format="mixed").dt.date
    b["CONDUCTOR"] = txt(col("CONDUCTOR"))
    b["LETRA"] = b["CONDUCTOR"].str.extract(r"(?i)\bruta\s+([A-Z]{1,2})\b")[0].str.upper().fillna("")
    b["ES_EASY"] = b["CLIENTE"].str.contains("EASY", na=False)
    b["ES_PARIS"] = b["CLIENTE"].str.contains("PARIS", na=False)
    b["ENTREGADO"] = b["ESTADO"].isin(ESTADOS_GESTIONADOS)
    b = b[b["PAT_KEY"].ne("") | b["ORDEN"].ne("")].reset_index(drop=True)
    b.attrs["mapeo"], b.attrs["avisos"] = mapa, avisos
    return b


def preparar_hela(df):
    mapa, avisos = mapear_columnas(df, HELA_CAMPOS)
    h = pd.DataFrame({k: _col(df, mapa, k) for k in HELA_CAMPOS}, index=df.index)
    h.attrs["mapeo"], h.attrs["avisos"] = mapa, avisos
    return h


def detalle_out(d):
    return pd.DataFrame({
        "Ruta": d["RUTA"], "Patente": d["PATENTE"], "Orden": d["ORDEN"], "Cliente": d["CLIENTE_TXT"],
        "Estado": d["ESTADO_TXT"], "Sub-estado": d["SUBESTADO"], "Usuario móvil": d["USUARIO"],
        "Fecha compromiso": d["F_COMP"].map(fmt_fecha), "Comuna": d["COMUNA"],
    }).reset_index(drop=True)


# =====================================================================
# 1. NS DIARIO
# =====================================================================
def procesar_ns(b, fecha, manual):
    f = b[b["ES_EASY"] | b["ES_PARIS"]]
    easy = f[f["ES_EASY"]]
    paris = f[f["ES_PARIS"]]

    def moviles(d):
        return int(d.loc[d["PAT_KEY"].ne(""), "PAT_KEY"].nunique())

    def fc(d):
        x = d[d["F_COMP"] == fecha]
        return int(len(x)), int(x["ENTREGADO"].sum())

    fce_t, fce_e = fc(easy)
    fcp_t, fcp_e = fc(paris)
    meta = dict(manual)
    meta.update({
        "fecha": fecha.isoformat(),
        "acid_easy_tot": int(len(easy)), "acid_easy_ent": int(easy["ENTREGADO"].sum()),
        "acid_paris_tot": int(len(paris)), "acid_paris_ent": int(paris["ENTREGADO"].sum()),
        "moviles_gen": moviles(f), "moviles_easy": moviles(easy), "moviles_paris": moviles(paris),
        "fc_btk_easy_tot": fce_t, "fc_btk_easy_ent": fce_e,
        "fc_btk_paris_tot": fcp_t, "fc_btk_paris_ent": fcp_e,
        "mapeo_btk": mapa_a_filas(b.attrs.get("mapeo", {})), "avisos_btk": list(b.attrs.get("avisos", [])),
    })
    noent = f[~f["ENTREGADO"]]
    sub = noent["SUBESTADO"].replace("", "(sin sub-estado)").value_counts().reset_index()
    sub.columns = ["Sub-estado No Entrega", "Cantidad"]
    dfs = {"subestados": sub, "detalle_no_entregados": detalle_out(noent), "detalle": detalle_out(f)}
    return meta, dfs


def calcular_ns(m):
    """Totales, no entregados y porcentajes de cada bloque, con las mismas fórmulas en todo el sistema."""
    def g(k):
        return int(m.get(k, 0) or 0)

    r = {}
    for bloque in ("fc", "acid", "ret"):
        t = {c: g(f"{bloque}_{c}_tot") for c in ("easy", "paris")}
        e = {c: g(f"{bloque}_{c}_ent") for c in ("easy", "paris")}
        t["gen"], e["gen"] = t["easy"] + t["paris"], e["easy"] + e["paris"]
        for c in ("easy", "paris", "gen"):
            r[f"{bloque}_{c}_tot"] = t[c]
            r[f"{bloque}_{c}_ent"] = e[c]
            r[f"{bloque}_{c}_no"] = t[c] - e[c]
            r[f"ns_{bloque}_{c}"] = pct(e[c], t[c])
    for k in ("moviles_gen", "moviles_easy", "moviles_paris"):
        r[k] = g(k)
    return r


def resumen_ns_df(m):
    c = calcular_ns(m)
    filas = []
    bloques = [("Fecha compromiso (FC)", "fc"), ("NS Ácido", "acid"), ("Retiros / recogidos", "ret")]
    for titulo, b in bloques:
        for nombre, k in (("Easy", "easy"), ("París", "paris"), ("General", "gen")):
            filas.append({"Sección": titulo, "Ámbito": nombre, "Total": c[f"{b}_{k}_tot"],
                          "Entregados": c[f"{b}_{k}_ent"], "No entregados": c[f"{b}_{k}_no"],
                          "NS (%)": round(c[f"ns_{b}_{k}"], 2)})
    for nombre, k in (("General", "moviles_gen"), ("Easy", "moviles_easy"), ("París", "moviles_paris")):
        filas.append({"Sección": "Rutas y móviles", "Ámbito": nombre, "Total": c[k],
                      "Entregados": None, "No entregados": None, "NS (%)": None})
    return pd.DataFrame(filas)


def pdf_ns(m, sub_df):
    """Informe PDF del NS diario, con el mismo formato del informe de ejemplo."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    c = calcular_ns(m)
    fecha = datetime.date.fromisoformat(m["fecha"]).strftime("%d/%m/%Y")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title=f"Informe de nivel de servicio {fecha}", author="Transportes Yáñez")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], alignment=0, fontSize=16, spaceAfter=4)
    seccion = ParagraphStyle("seccion", parent=estilos["Heading2"], fontSize=12, spaceBefore=8, spaceAfter=4)

    def tabla(datos, negrita_ultima=False):
        t = Table(datos, repeatRows=1, hAlign="LEFT")
        est = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EDF2")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C0CC")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ]
        if negrita_ultima:
            est.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        t.setStyle(TableStyle(est))
        return t

    historia = [Paragraph(f"Informe de nivel de servicio | {fecha}", titulo)]

    historia.append(Paragraph("1. Fecha compromiso (FC)", seccion))
    datos = [["Cliente", "Total FC", "Entregados", "No entregados", "Cumplimiento"]]
    for nombre, k in (("Easy", "easy"), ("París", "paris"), ("General (calculado)", "gen")):
        datos.append([nombre, str(c[f"fc_{k}_tot"]), str(c[f"fc_{k}_ent"]), str(c[f"fc_{k}_no"]),
                      fmt_pct(c[f"ns_fc_{k}"])])
    historia.append(tabla(datos))

    historia.append(Paragraph("2. NS Ácido", seccion))
    datos = [["Ámbito", "Total", "Entregados", "NS Ácido"]]
    for nombre, k in (("General", "gen"), ("Easy", "easy"), ("París", "paris")):
        datos.append([nombre, str(c[f"acid_{k}_tot"]), str(c[f"acid_{k}_ent"]), fmt_pct(c[f"ns_acid_{k}"])])
    historia.append(tabla(datos))

    historia.append(Paragraph("3. Retiros / recogidos", seccion))
    datos = [["Cliente", "Retiros totales", "Retiros entregados", "Tasa calculada"]]
    for nombre, k in (("Easy", "easy"), ("París", "paris"), ("General", "gen")):
        datos.append([nombre, str(c[f"ret_{k}_tot"]), str(c[f"ret_{k}_ent"]), fmt_pct(c[f"ns_ret_{k}"])])
    historia.append(tabla(datos))

    historia.append(Paragraph("4. Rutas y móviles", seccion))
    datos = [["Indicador", "Cantidad"],
             ["Rutas totales móviles", str(c["moviles_gen"])],
             ["Easy", str(c["moviles_easy"])],
             ["París", str(c["moviles_paris"])]]
    historia.append(tabla(datos))

    historia.append(Paragraph("5. Sub-estados de no entrega del día", seccion))
    datos = [["Sub-estado", "Cantidad"]]
    total = 0
    if sub_df is not None and not sub_df.empty:
        for _, fila in sub_df.iterrows():
            datos.append([str(fila.iloc[0]), str(int(fila.iloc[1]))])
            total += int(fila.iloc[1])
    else:
        datos.append(["Sin no entregas", "0"])
    datos.append(["Total de registros listados (calculado)", str(total)])
    historia.append(tabla(datos, negrita_ultima=True))
    historia.append(Spacer(1, 6))

    doc.build(historia)
    return buf.getvalue()


# =====================================================================
# 2. COMUNAS
# =====================================================================
ORIGEN_OK = "Hela + BTK"
ORIGEN_SOLO_BTK = "Solo BTK (no estaba en Hela)"
ORIGEN_SOLO_HELA = "Solo Hela (sin ruta en BTK)"


def _partes_comunas(texto):
    texto = "" if texto is None or pd.isna(texto) else str(texto)
    return [x.strip() for x in re.split(r"\s+-\s+|[,;/]", texto) if limpiar_texto(x)]


def _clave_comuna(nombre):
    r = obtener_tarifa_comuna(nombre)
    return r[0] if r else limpiar_texto(nombre)


def procesar_comunas(hela, b, fecha, nombres_pat=None):
    """Cruza Hela con BTK. El BTK manda. Cada ruta tiene una sola patente.
    - Si el Hela trae Patente, cada fila se une con el BTK por patente.
    - Si no la trae, cada fila del Hela es una ruta (A, B, C… en el orden del archivo). Se une con el BTK por
      la letra de la columna «conductor» ("RUTA A - NOMBRE"): la patente de la ruta es la que llevó más pedidos
      con esa letra. Desde ahí, puntos, entregados, compromisos, comunas y valor salen de TODOS los pedidos que
      llevó esa patente, aunque algunos vengan con la letra de otra ruta.
    La comuna pagada sale de los pedidos del BTK; las comunas que no estaban en Hela se agregan.
    nombres_pat (opcional, ver nombres_habituales): si viene, el driver que se muestra es el de la patente que hizo
    la ruta ese día, no el que traiga una etiqueta suelta de un pedido que cambió de ruta.
    Solo se consideran pedidos de clientes Easy y París."""
    attrs_b = dict(b.attrs)
    b = b[b["ES_EASY"] | b["ES_PARIS"]]
    h = preparar_hela(hela)
    por_patente = "PATENTE" in h.attrs["mapeo"]
    avisos_hela = list(h.attrs["avisos"])
    filas = []

    def fila_ruta(bp, patente, driver, nombre_ruta, comunas_hela, pts_hela, comp_hela, origen,
                  letra="", otras_rutas="", otras_patentes="", n_letra=None):
        easy = bp[bp["ES_EASY"]]
        paris = bp[bp["ES_PARIS"]]
        tot = int(len(bp))
        ent = int(bp["ENTREGADO"].sum())
        ent_easy = int(easy["ENTREGADO"].sum())
        ent_paris = int(paris["ENTREGADO"].sum())
        pts = pd.to_numeric(pts_hela, errors="coerce")
        comp = pd.to_numeric(comp_hela, errors="coerce")
        comp_btk = int((bp["F_COMP"] == fecha).sum())

        usuarios = bp.loc[bp["USUARIO"].ne(""), "USUARIO"]
        usuario_btk = usuarios.iloc[0] if len(usuarios) else ""
        if driver is None or pd.isna(driver) or not str(driver).strip():
            driver = usuario_btk

        # Comunas: las del Hela + las que aparecen en el BTK y no estaban en Hela
        partes_hela = list({_clave_comuna(x): x for x in reversed(_partes_comunas(comunas_hela))}.values())[::-1]
        claves_h = {_clave_comuna(x) for x in partes_hela}
        comunas_btk = list(dict.fromkeys(bp.loc[bp["COMUNA"].ne(""), "COMUNA"]))
        extras = [c for c in comunas_btk if _clave_comuna(c) not in claves_h]
        comunas_txt = " - ".join(partes_hela + extras)

        nombre_h, valor_h, desc_h = comuna_mas_lejana(" - ".join(partes_hela))
        nombre_b, valor_b, desc_b = ("SIN COMUNA RECONOCIDA", 0, [])
        if comunas_btk:
            nombre_b, valor_b, desc_b = comuna_mas_lejana(",".join(comunas_btk))
        if comunas_btk and nombre_b != "SIN COMUNA RECONOCIDA":
            pagada, valor = nombre_b, valor_b
        else:
            pagada, valor = nombre_h, valor_h
        desconocidas = list(dict.fromkeys(desc_h + desc_b))
        coincide_comuna = ("SÍ" if pagada == nombre_h else "NO") if (tot and partes_hela) else ""

        if origen == ORIGEN_SOLO_BTK:
            s_pts = s_comp = "SIN HELA"
        elif tot == 0:
            s_pts = s_comp = "SIN BTK"
        else:
            s_pts = "SÍ" if pts == tot else "NO"
            s_comp = "SÍ" if comp == comp_btk else "NO"
        dif_pts = int(tot - pts) if origen == ORIGEN_OK and tot and not pd.isna(pts) else None
        dif_comp = int(comp_btk - comp) if origen == ORIGEN_OK and tot and not pd.isna(comp) else None

        if tot:
            rutas = bp["RUTA"][bp["RUTA"].ne("")]
            ruta_btk = rutas.mode().iloc[0] if len(rutas) else ""
        else:
            ruta_btk = "SIN RUTA BTK"

        return {
            "A: Fecha": fecha.strftime("%d/%m/%Y"),
            "B: Driver": driver,
            "C: Patente": patente,
            "D: Nombre ruta": nombre_ruta,
            "E: Comunas": comunas_txt,
            "F: Ruta BTK": ruta_btk,
            "G: Total puntos HELA": None if pd.isna(pts) else int(pts),
            "H: Compromisos HELA": None if pd.isna(comp) else int(comp),
            "I: Comuna Lejana": pagada,
            "J: Valor": valor,
            "K: Entregados BTK": ent,
            "L: No Entregados BTK": tot - ent,
            "M: Entregados Easy BTK": ent_easy,
            "N: No Entregados Easy BTK": int(len(easy)) - ent_easy,
            "O: Entregados París BTK": ent_paris,
            "P: No Entregados París BTK": int(len(paris)) - ent_paris,
            "Q: NS Patente (%)": round(pct(ent, tot), 2),
            "R: Puntos BTK": tot,
            "S: Pedidos Easy": int(len(easy)),
            "T: Pedidos París": int(len(paris)),
            "U: Coincide puntos": s_pts,
            "V: Compromisos BTK": comp_btk,
            "W: Coincide compromisos": s_comp,
            "X: Usuario móvil BTK": usuario_btk,
            "Y: Comuna lejana según Hela": nombre_h if partes_hela else "",
            "Z: Coincide con Hela": coincide_comuna,
            "AA: Comunas no reconocidas": ", ".join(desconocidas),
            "AB: Origen": origen,
            "AC: Dif. puntos (BTK − Hela)": dif_pts,
            "AD: Dif. compromisos (BTK − Hela)": dif_comp,
            "AE: Comunas agregadas desde BTK": " - ".join(extras) if partes_hela else "",
            "AF: Valor según Hela": valor_h if partes_hela else None,
            "AG: Letra de ruta": letra,
            "AH: Pedidos de otras rutas en esta patente": otras_rutas,
            "AI: Pedidos de esta ruta llevados por otra patente": otras_patentes,
            "AJ: Pedidos con la letra de esta ruta (BTK)": n_letra,
        }

    def etiqueta_conductor(bp, letra=None):
        if nombres_pat and len(bp):
            nombre = nombre_del_dia(nombres_pat.get(bp["PAT_KEY"].iloc[0]), bp)
            if nombre:
                return f"RUTA {letra} - {nombre}" if letra else nombre
        c = bp["CONDUCTOR"][bp["CONDUCTOR"].ne("")]
        if letra:
            c = bp.loc[bp["LETRA"].eq(letra) & bp["CONDUCTOR"].ne(""), "CONDUCTOR"]
        return c.value_counts().index[0] if len(c) else ""

    solo_btk = []
    avisos_rutas = []
    if por_patente:
        entradas, por_clave, sin_patente, agrupadas = [], {}, 0, []
        k = 0

        def suma(a, b_):
            a, b_ = pd.to_numeric(a, errors="coerce"), pd.to_numeric(b_, errors="coerce")
            return b_ if pd.isna(a) else (a if pd.isna(b_) else a + b_)

        for _, row in h.iterrows():
            if (pd.isna(row["NOMBRE_RUTA"]) and pd.isna(row["COMUNAS"]) and pd.isna(row["PUNTOS"])
                    and pd.isna(row["PATENTE"])):
                continue
            letra = letra_excel(k)  # la ruta A, B, C… sigue el orden de las filas del Hela
            k += 1
            pat_raw = row["PATENTE"]
            if pd.isna(pat_raw) or not str(pat_raw).strip():
                sin_patente += 1
                entradas.append({"clave": None, "letras": [letra], "driver": row["DRIVER"], "nombres": [row["NOMBRE_RUTA"]],
                                 "comunas": [row["COMUNAS"]], "puntos": row["PUNTOS"], "comp": row["COMPROMISOS"],
                                 "patente": ""})
                continue
            patente = str(pat_raw).upper().strip()
            clave = clave_patente(patente)
            if clave in por_clave:
                # una patente con más de una ruta el mismo día (p. ej. una «ruta adicional»): se suman en una sola fila
                e_ = por_clave[clave]
                e_["letras"].append(letra)
                e_["nombres"].append(row["NOMBRE_RUTA"])
                e_["comunas"].append(row["COMUNAS"])
                e_["puntos"] = suma(e_["puntos"], row["PUNTOS"])
                e_["comp"] = suma(e_["comp"], row["COMPROMISOS"])
                agrupadas.append(f"{patente} ({' + '.join(e_['letras'])})")
            else:
                por_clave[clave] = {"clave": clave, "letras": [letra], "driver": row["DRIVER"],
                                    "nombres": [row["NOMBRE_RUTA"]], "comunas": [row["COMUNAS"]],
                                    "puntos": row["PUNTOS"], "comp": row["COMPROMISOS"], "patente": patente}
                entradas.append(por_clave[clave])
        claves_hela = set(por_clave)
        for e_ in entradas:
            bp = b[b["PAT_KEY"] == e_["clave"]] if e_["clave"] else b.iloc[0:0]
            origen = ORIGEN_OK if len(bp) else ORIGEN_SOLO_HELA
            patente = bp["PATENTE"].iloc[0] if len(bp) else e_["patente"]  # como viene en el BTK
            letra = " + ".join(e_["letras"])
            driver = e_["driver"]
            if (driver is None or pd.isna(driver) or not str(driver).strip()) and len(bp):
                driver = etiqueta_conductor(bp, e_["letras"][0])
            comunas = " - ".join(str(c_) for c_ in e_["comunas"] if not pd.isna(c_))
            nombres = " + ".join(str(n_) for n_ in e_["nombres"] if not pd.isna(n_))
            n_letra = int(b["LETRA"].isin(e_["letras"]).sum())
            filas.append(fila_ruta(bp, patente, driver, nombres, comunas, e_["puntos"], e_["comp"], origen, letra,
                                   n_letra=n_letra))
        for clave, g in b[b["PAT_KEY"].ne("")].groupby("PAT_KEY", sort=False):
            if clave in claves_hela:
                continue
            patente = g["PATENTE"].iloc[0]
            solo_btk.append(patente)
            filas.append(fila_ruta(g, patente, etiqueta_conductor(g), "(solo en BTK)", "", None, None,
                                   ORIGEN_SOLO_BTK))
        if sin_patente:
            avisos_rutas.append(f"{sin_patente} ruta(s) del Hela sin patente (quedan como «Solo Hela»).")
        if agrupadas:
            avisos_rutas.append("Patente con más de una ruta el mismo día (se sumaron en una sola fila): "
                                + ", ".join(dict.fromkeys(agrupadas)))
    else:
        avisos_hela = [a for a in avisos_hela if "Patente" not in a and "Driver" not in a]
        if len(b) and not (b["LETRA"] != "").any():
            raise ValueError("El Hela no trae patentes y el BTK no trae la letra de ruta en la columna «conductor» "
                             "(por ejemplo «RUTA A - NOMBRE»). No hay forma de unir ambos archivos.")
        rutas_hela = []
        for _, row in h.iterrows():
            if pd.isna(row["NOMBRE_RUTA"]) and pd.isna(row["COMUNAS"]) and pd.isna(row["PUNTOS"]):
                continue
            rutas_hela.append((letra_excel(len(rutas_hela)), row))

        # Una patente por ruta: se empareja letra <-> patente empezando por la combinación con más pedidos
        cuenta = (b[b["LETRA"].ne("") & b["PAT_KEY"].ne("")].groupby(["LETRA", "PAT_KEY"]).size()
                  .reset_index(name="n").sort_values(["n", "LETRA"], ascending=[False, True]))
        pat_de_letra, letra_de_pat = {}, {}
        for _, r in cuenta.iterrows():
            if r["LETRA"] in pat_de_letra or r["PAT_KEY"] in letra_de_pat:
                continue
            pat_de_letra[r["LETRA"]] = r["PAT_KEY"]
            letra_de_pat[r["PAT_KEY"]] = r["LETRA"]

        def cruces(bp, letra):
            otras = bp[bp["LETRA"] != letra]["LETRA"].replace("", "sin letra").value_counts()
            return " · ".join(f"{k}: {v}" for k, v in otras.items())

        usados = set()
        for letra, row in rutas_hela:
            pk = pat_de_letra.get(letra)
            bp = b[b["PAT_KEY"] == pk] if pk else b.iloc[0:0]
            n_letra = int((b["LETRA"] == letra).sum())
            if pk:
                usados.add(pk)
                ajenas = b[(b["LETRA"] == letra) & (b["PAT_KEY"] != pk)]
                otras_pat = " · ".join(f"{k}: {v}" for k, v in ajenas["PATENTE"].value_counts().items())
                filas.append(fila_ruta(bp, bp["PATENTE"].iloc[0], etiqueta_conductor(bp, letra), row["NOMBRE_RUTA"],
                                       row["COMUNAS"], row["PUNTOS"], row["COMPROMISOS"], ORIGEN_OK, letra,
                                       cruces(bp, letra), otras_pat, n_letra))
            else:
                filas.append(fila_ruta(bp, "", "", row["NOMBRE_RUTA"], row["COMUNAS"], row["PUNTOS"],
                                       row["COMPROMISOS"], ORIGEN_SOLO_HELA, letra, "", "", n_letra))

        # Patentes que salieron a ruta y no quedaron emparejadas con una ruta del Hela
        for clave, g in b[b["PAT_KEY"].ne("")].groupby("PAT_KEY", sort=False):
            if clave in usados:
                continue
            letra = letra_de_pat.get(clave, "")
            solo_btk.append(g["PATENTE"].iloc[0])
            nombre = "(ruta no está en Hela)" if letra else "(sin letra de ruta en BTK)"
            filas.append(fila_ruta(g, g["PATENTE"].iloc[0], etiqueta_conductor(g, letra or None), nombre, "",
                                   None, None, ORIGEN_SOLO_BTK, letra, cruces(g, letra), "",
                                   int((b["LETRA"] == letra).sum()) if letra else None))

    df = pd.DataFrame(filas)
    info = {"modo_union": "patente" if por_patente else "letra",
            "patentes_solo_btk": solo_btk,
            "patentes_solo_hela": df.loc[df["AB: Origen"] == ORIGEN_SOLO_HELA, "C: Patente"].tolist() if len(df) else [],
            "mapeo_hela": mapa_a_filas(h.attrs["mapeo"]), "avisos_hela": avisos_hela + avisos_rutas,
            "avisos_rutas": avisos_rutas,
            "mapeo_btk": mapa_a_filas(attrs_b.get("mapeo", {})), "avisos_btk": list(attrs_b.get("avisos", []))}
    return df, info


def fecha_desde_nombre(nombre, hoy):
    """Lee dd-mm (o dd-mm-aaaa) al final del nombre del archivo, p. ej. «…Paris29-09.xlsx»."""
    m = re.search(r"(?<!\d)(\d{2})[-_.](\d{2})(?:[-_.](\d{2}|\d{4}))?\.[A-Za-z0-9]+$", str(nombre))
    if not m:
        return None
    dia, mes, anio = int(m.group(1)), int(m.group(2)), m.group(3)
    try:
        if anio:
            return datetime.date(int(anio) + (2000 if len(anio) == 2 else 0), mes, dia)
        f = datetime.date(hoy.year, mes, dia)
        return f if f <= hoy + datetime.timedelta(days=30) else datetime.date(hoy.year - 1, mes, dia)
    except ValueError:
        return None


def fecha_detectada(tipo, df, nombre=None, hoy=None):
    """Fecha del archivo: BTK Yáñez -> columna «Fecha ruta»; si no hay, la del nombre del archivo.
    (La fecha dentro del nombre de ruta del Hela es la del día anterior, por eso no se usa.)"""
    try:
        if tipo == "btk_yanez":
            cols = {limpiar_texto(c): c for c in df.columns}
            if "FECHA RUTA" in cols:
                f = pd.to_datetime(df[cols["FECHA RUTA"]], errors="coerce", format="mixed",
                                   dayfirst=True).dt.date.dropna()
                if len(f):
                    return f.mode().iloc[0]
    except Exception:  # noqa: BLE001
        pass
    return fecha_desde_nombre(nombre, hoy or hoy_chile()) if nombre else None


def df_bonificacion(dfc, excluir_sin_ruta=True, incluir_cero=False):
    """Una fila por ruta que tuvo pago. Las rutas con valor $0 (no llevaron ninguna comuna con valor) no se agregan."""
    d = dfc.copy()
    if excluir_sin_ruta:
        d = d[d["F: Ruta BTK"] != "SIN RUTA BTK"]
    if not incluir_cero:
        d = d[pd.to_numeric(d["J: Valor"], errors="coerce").fillna(0) > 0]
    return pd.DataFrame({
        "A: Fecha": d["A: Fecha"], "B: Patente": d["C: Patente"], "C: Ruta": d["F: Ruta BTK"],
        "D: Valor": d["J: Valor"], "E: Fecha (repetida)": d["A: Fecha"],
        "F: Comuna pagada": d["I: Comuna Lejana"],
    }).reset_index(drop=True)


def bonificacion_por_patente(dfb):
    if dfb.empty:
        return pd.DataFrame(columns=["Patente", "Días", "Total bonificación"])
    g = dfb.groupby("B: Patente").agg(**{"Días": ("A: Fecha", "nunique"),
                                         "Total bonificación": ("D: Valor", "sum")}).reset_index()
    g = g.rename(columns={"B: Patente": "Patente"})
    return g.sort_values("Total bonificación", ascending=False).reset_index(drop=True)


# =====================================================================
# 4. CALCES PARÍS
# =====================================================================
def _subestados_equivalentes(sub_y, sub_p):
    for a, b in REGLAS_CALCE:
        if a in sub_y and b in sub_p:
            return True
    if not sub_y:
        return False
    # París suele anteponer el estado: "No Recogido - Cliente No Está" contiene "Cliente No Está"
    return sub_y == sub_p or sub_y in sub_p


def evaluar_calce(est_y, sub_y, est_p, sub_p, encontrado):
    """Todos los textos vienen ya limpios (mayúsculas, sin tildes). Devuelve (coincide, observación)."""
    if not encontrado:
        return False, "Guía no encontrada en BTK París"
    if est_p not in ESTADOS_CONOCIDOS:
        return False, f"París aún no la calza (su estado es «{est_p.capitalize() or 'vacío'}»)"
    # Retiros: el estado es Recogido / No recogido en ambos sistemas (los sub-estados derivan)
    if est_y == "RECOGIDO":
        return (True, "") if est_p == "RECOGIDO" else (False, "Yáñez está Recogido y París no")
    if est_y == "NO RECOGIDO":
        if est_p != "NO RECOGIDO":
            return False, "Yáñez está No recogido y París no"
        if _subestados_equivalentes(sub_y, sub_p):
            return True, ""
        return False, "Sub-estados no equivalentes"
    # Despachos
    if est_y == "ENTREGADO" and "EN CLIENTE" in sub_p:
        return True, ""
    if _subestados_equivalentes(sub_y, sub_p):
        return True, ""
    conocido = est_y == "ENTREGADO" or any(a in sub_y for a, _ in REGLAS_CALCE)
    return False, ("Sub-estados no equivalentes" if conocido else "Sin regla definida para este sub-estado")


def procesar_calces(y, p_raw, fecha):
    mapa_p, avisos_p = mapear_columnas(p_raw, PARIS_CAMPOS)
    yp = y[y["ES_PARIS"]].copy()
    p = pd.DataFrame({
        "ORDEN": txt(_col(p_raw, mapa_p, "ORDEN")),
        "EST_P": txt(_col(p_raw, mapa_p, "ESTADO")),
        "SUB_P": txt(_col(p_raw, mapa_p, "SUBESTADO")),
    })
    duplicadas = int(p["ORDEN"].duplicated().sum())
    p = p.drop_duplicates("ORDEN", keep="first")
    m = yp.merge(p, on="ORDEN", how="left")
    m["ENCONTRADO"] = m["EST_P"].notna()
    m["EST_P_C"] = m["EST_P"].fillna("").map(limpiar_texto)
    m["SUB_P_C"] = m["SUB_P"].fillna("").map(limpiar_texto)

    res = [evaluar_calce(e, s, ep, sp, enc) for e, s, ep, sp, enc in
           zip(m["ESTADO"], m["SUB_CLEAN"], m["EST_P_C"], m["SUB_P_C"], m["ENCONTRADO"])]
    m["OK"] = [r[0] for r in res]
    m["OBS"] = [r[1] for r in res]

    df = pd.DataFrame({
        "Fecha": fecha.strftime("%d/%m/%Y"),
        "N° de pedido": m["ORDEN"],
        "Beetrack T.Yañez Estado": m["ESTADO_TXT"],
        "Beetrack T.Yañez Sub Estado": m["SUBESTADO"],
        "Beetrack París": m["EST_P"].fillna("NO ENCONTRADO"),
        "Beetrack París Sub Estado": m["SUB_P"].fillna("NO ENCONTRADO"),
        "Estado Cruce": m["OK"].map({True: "Coincide", False: "No Coincide"}),
        "Observación": "",  # para tus notas (p. ej. "Se solicita calce") al bajar el Excel
        "Fecha Compromiso": m["F_COMP"].map(fmt_fecha),
        "Diagnóstico automático": m["OBS"],
    })
    info = {"guias_duplicadas_paris": duplicadas,
            "mapeo_paris": mapa_a_filas(mapa_p), "avisos_paris": list(avisos_p),
            "mapeo_btk": mapa_a_filas(y.attrs.get("mapeo", {})), "avisos_btk": list(y.attrs.get("avisos", []))}
    return df.reset_index(drop=True), info


# =====================================================================
# 5. SEMANAL / CONSOLIDADOS
# =====================================================================
def tabla_fc(datos, texto=False):
    """Tabla con una columna por fecha: NS %, pedidos totales y no entregados, más el total del período."""
    calcs = [(f, calcular_ns(m)) for f, m, _ in datos]
    fp = fmt_pct if texto else (lambda x: round(x, 2))
    fi = (lambda x: str(int(x))) if texto else int
    fechas = [f.strftime("%d/%m/%Y") for f, _ in calcs]

    def S(k):
        return sum(c[k] for _, c in calcs)

    filas = [
        ("NS Fecha Compromiso General (%)", lambda c: c["ns_fc_gen"], pct(S("fc_gen_ent"), S("fc_gen_tot")), fp),
        ("NS FC Easy (%)", lambda c: c["ns_fc_easy"], pct(S("fc_easy_ent"), S("fc_easy_tot")), fp),
        ("NS FC París (%)", lambda c: c["ns_fc_paris"], pct(S("fc_paris_ent"), S("fc_paris_tot")), fp),
        ("Pedidos totales FC", lambda c: c["fc_gen_tot"], S("fc_gen_tot"), fi),
        ("Entregados FC", lambda c: c["fc_gen_ent"], S("fc_gen_ent"), fi),
        ("No entregados FC", lambda c: c["fc_gen_no"], S("fc_gen_no"), fi),
        ("% No entregados FC", lambda c: pct(c["fc_gen_no"], c["fc_gen_tot"]),
         pct(S("fc_gen_no"), S("fc_gen_tot")), fp),
    ]
    out = []
    for etiqueta, f_dia, total, fmt in filas:
        fila = {"Indicador": etiqueta}
        for (_, c), ftxt in zip(calcs, fechas):
            fila[ftxt] = fmt(f_dia(c))
        fila["TOTAL PERÍODO"] = fmt(total)
        out.append(fila)
    return pd.DataFrame(out)


def df_ns_dias(datos):
    filas = []
    for f, m, _ in datos:
        c = calcular_ns(m)
        filas.append({
            "Fecha": f.strftime("%d/%m/%Y"),
            "FC totales": c["fc_gen_tot"], "FC entregados": c["fc_gen_ent"], "FC no entregados": c["fc_gen_no"],
            "NS FC General (%)": round(c["ns_fc_gen"], 2), "NS FC Easy (%)": round(c["ns_fc_easy"], 2),
            "NS FC París (%)": round(c["ns_fc_paris"], 2),
            "Órdenes a ruta": c["acid_gen_tot"], "Entregadas": c["acid_gen_ent"], "No entregadas": c["acid_gen_no"],
            "NS Ácido General (%)": round(c["ns_acid_gen"], 2), "NS Ácido Easy (%)": round(c["ns_acid_easy"], 2),
            "NS Ácido París (%)": round(c["ns_acid_paris"], 2),
            "Retiros totales": c["ret_gen_tot"], "Retiros entregados": c["ret_gen_ent"],
            "NS Retiros General (%)": round(c["ns_ret_gen"], 2), "NS Retiros Easy (%)": round(c["ns_ret_easy"], 2),
            "NS Retiros París (%)": round(c["ns_ret_paris"], 2),
            "Móviles general": c["moviles_gen"], "Móviles Easy": c["moviles_easy"], "Móviles París": c["moviles_paris"],
        })
    return pd.DataFrame(filas)


def resumen_periodo(datos):
    """Métricas globales del rango, sumando todos los días."""
    calcs = [calcular_ns(m) for _, m, _ in datos]

    def S(k):
        return sum(c[k] for c in calcs)

    return {
        "ordenes": S("acid_gen_tot"), "entregadas": S("acid_gen_ent"), "no_entregadas": S("acid_gen_no"),
        "ns_acid_gen": pct(S("acid_gen_ent"), S("acid_gen_tot")),
        "ns_acid_easy": pct(S("acid_easy_ent"), S("acid_easy_tot")),
        "ns_acid_paris": pct(S("acid_paris_ent"), S("acid_paris_tot")),
        "pct_no_ent": pct(S("acid_gen_no"), S("acid_gen_tot")),
        "prom_moviles": S("moviles_gen") / len(calcs) if calcs else 0,
        "rutas_easy": S("moviles_easy"), "rutas_paris": S("moviles_paris"),
    }


def ranking_submotivos(datos):
    dfs = [d["subestados"] for _, _, d in datos if "subestados" in d and not d["subestados"].empty]
    if not dfs:
        return pd.DataFrame(columns=["Causal / Submotivo No Entrega", "Total Impacto Acumulado"])
    r = pd.concat(dfs, ignore_index=True)
    r.columns = ["Causal / Submotivo No Entrega", "Total Impacto Acumulado"]
    r = r.groupby("Causal / Submotivo No Entrega", as_index=False)["Total Impacto Acumulado"].sum()
    return r.sort_values("Total Impacto Acumulado", ascending=False).reset_index(drop=True)


def concat_dfs(datos, nombre, con_fecha=False):
    piezas = []
    for f, _, dfs in datos:
        d = dfs.get(nombre)
        if d is None or d.empty:
            continue
        if con_fecha:
            d = d.copy()
            d.insert(0, "Fecha", f.strftime("%d/%m/%Y"))
        piezas.append(d)
    return pd.concat(piezas, ignore_index=True) if piezas else pd.DataFrame()


# =====================================================================
# EXCEL
# =====================================================================
def _es_pct(nombre):
    t = str(nombre)
    return "(%)" in t or t.startswith("%")


def _con_porcentajes(df):
    """Excel guarda los porcentajes como fracción (0,9231) con formato de %, para que se puedan formatear y operar.
    Devuelve (df, columnas_en_porcentaje, posiciones_de_filas_en_porcentaje)."""
    d = df.copy()
    cols = [c for c in d.columns if _es_pct(c) and pd.api.types.is_numeric_dtype(d[c])]
    for c in cols:
        d[c] = (d[c] / 100).round(6)
    filas = []
    if "Indicador" in d.columns:
        otras = [c for c in d.columns if c != "Indicador"]
        d[otras] = d[otras].astype(object)
        for pos, (_, fila) in enumerate(d.iterrows()):
            if _es_pct(fila["Indicador"]):
                filas.append(pos)
                for c in otras:
                    v = d.iloc[pos][c]
                    if isinstance(v, (int, float)) and not pd.isna(v):
                        d.iloc[pos, d.columns.get_loc(c)] = round(v / 100, 6)
    return d, cols, filas


def a_excel(hojas, graficos=None):
    """hojas = {'Nombre': DataFrame}. Encabezado en negrita, filtros, primera fila fija, ancho automático y
    los porcentajes con formato de % (las columnas con «(%)» o que empiezan con «%»).
    graficos = {"Nombre de hoja": {"titulo", "x": columna de fechas, "series": [columnas], "colores": ["RRGGBB"]}}
    agrega un gráfico de líneas a la derecha de la tabla."""
    from openpyxl.styles import Font, PatternFill

    if not hojas:
        hojas = {"Sin datos": pd.DataFrame({"Aviso": ["No hay datos para exportar"]})}
    buf = io.BytesIO()
    usados = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            n = re.sub(r"[\[\]\:\*\?\/\\]", "-", nombre)[:31]
            base, i = n, 2
            while n in usados:
                n = f"{base[:28]}_{i}"
                i += 1
            usados.add(n)
            d, cols_pct, filas_pct = _con_porcentajes(df)
            d.to_excel(writer, sheet_name=n, index=False)
            ws = writer.sheets[n]
            ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for celda in ws[1]:
                celda.font = Font(bold=True)
                celda.fill = PatternFill("solid", fgColor="E9EDF2")
            for c in cols_pct:
                col_n = list(d.columns).index(c) + 1
                for fila in range(2, ws.max_row + 1):
                    ws.cell(fila, col_n).number_format = "0.00%"
            for pos in filas_pct:
                for col_n in range(2, ws.max_column + 1):
                    ws.cell(pos + 2, col_n).number_format = "0.00%"
            for col in ws.columns:
                largo = max((len(str(c.value)) for c in col if c.value is not None), default=0)
                ws.column_dimensions[col[0].column_letter].width = min(largo + 2, 45)
            if graficos and nombre in graficos and ws.max_row > 1:
                _grafico_lineas_excel(ws, d, graficos[nombre])
    return buf.getvalue()


def _grafico_lineas_excel(ws, d, g):
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter

    cols = list(d.columns)
    ch = LineChart()
    ch.title = g.get("titulo", "")
    ch.height, ch.width = 9.5, 26
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = "0%"
    vals = [v for c in g["series"] for v in d[c].dropna()]
    if vals:
        ch.y_axis.scaling.min = max(0.0, (int(min(vals) * 100 - 5) // 10) / 10)
        ch.y_axis.scaling.max = 1.0
    for i, c in enumerate(g["series"]):
        ch.add_data(Reference(ws, min_col=cols.index(c) + 1, min_row=1, max_row=ws.max_row), titles_from_data=True)
        serie = ch.series[i]
        serie.smooth = False
        if i < len(g.get("colores", [])):
            serie.graphicalProperties.line.solidFill = g["colores"][i]
        serie.graphicalProperties.line.width = 19000
    ch.set_categories(Reference(ws, min_col=cols.index(g["x"]) + 1, min_row=2, max_row=ws.max_row))
    ws.add_chart(ch, f"{get_column_letter(ws.max_column + 2)}2")


GRAFICO_COMPROMISO_EXCEL = {"NS Compromiso día": {"titulo": "NS Compromiso por día", "x": "Fecha",
                                                  "series": ["NS Easy (%)", "NS París (%)", "NS Compromiso (%)"],
                                                  "colores": ["1F5FBF", "7FC4FF", "FF2B2B"]}}
GRAFICO_ACIDO_EXCEL = {"NS Ácido día": {"titulo": "NS Ácido por día", "x": "Fecha",
                                        "series": ["NS Easy (%)", "NS París (%)", "NS Ácido (%)"],
                                        "colores": ["1F5FBF", "7FC4FF", "FF2B2B"]}}


def excel_completo(f_ini, f_fin):
    hojas = {}
    ns = cargar("ns", f_ini, f_fin)
    if ns:
        hojas["NS por día"] = df_ns_dias(ns)
        hojas["NS FC por fecha"] = tabla_fc(ns)
        hojas["Sub-estados acumulados"] = ranking_submotivos(ns)
        det = concat_dfs(ns, "detalle_no_entregados", con_fecha=True)
        if not det.empty:
            hojas["No entregados detalle"] = det
    com = cargar("comunas", f_ini, f_fin)
    dfc = concat_dfs(com, "reporte")
    if not dfc.empty:
        hojas["Comunas"] = dfc
        bon = df_bonificacion(dfc)
        hojas["Bonificación"] = bon
        hojas["Bonificación por patente"] = bonificacion_por_patente(bon)
    cal = cargar("calces", f_ini, f_fin)
    dcal = concat_dfs(cal, "cruce")
    if not dcal.empty:
        hojas["Calces París"] = dcal
    return a_excel(hojas) if hojas else None


# =====================================================================
# 7. REPORTE MENSUAL
# =====================================================================
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]


def primer_dia_mes(anio, mes):
    return datetime.date(int(anio), int(mes), 1)


def inicio_semana(d):
    return d - datetime.timedelta(days=d.weekday())


def etiqueta_semana(fechas):
    return f"Semana del {min(fechas):%d/%m} al {max(fechas):%d/%m}"


def mes_detectado(df):
    """(año, mes) más frecuente de la columna de fecha del archivo, o None."""
    try:
        mapa, _ = mapear_columnas(df, {"FECHA": BTK_CAMPOS["FECHA"]})
        if "FECHA" not in mapa:
            return None
        f = pd.to_datetime(df.iloc[:, mapa["FECHA"]["indice"]], errors="coerce", dayfirst=True,
                           format="mixed").dropna()
        if f.empty:
            return None
        m = f.dt.to_period("M").mode().iloc[0]
        return m.year, m.month
    except Exception:  # noqa: BLE001
        return None


def _fila_acido(g):
    tot, ent = int(len(g)), int(g["ENTREGADO"].sum())
    easy, paris = g[g["ES_EASY"]], g[g["ES_PARIS"]]
    return {
        "Órdenes totales": tot, "Entregadas / recogidas": ent, "No entregadas": tot - ent,
        "NS Ácido (%)": round(pct(ent, tot), 2),
        "Easy totales": int(len(easy)), "Easy entregadas": int(easy["ENTREGADO"].sum()),
        "NS Easy (%)": round(pct(int(easy["ENTREGADO"].sum()), int(len(easy))), 2),
        "París totales": int(len(paris)), "París entregadas": int(paris["ENTREGADO"].sum()),
        "NS París (%)": round(pct(int(paris["ENTREGADO"].sum()), int(len(paris))), 2),
    }


def tablas_ns_acido(b):
    """NS ácido por día, por semana (lunes a domingo) y del mes: entregadas o recogidas / órdenes totales."""
    e = b[(b["ES_EASY"] | b["ES_PARIS"]) & b["FECHA"].notna()].copy()
    cols_dia = ["Fecha", "Semana"]
    if e.empty:
        vacio = pd.DataFrame(columns=cols_dia + list(_fila_acido(e).keys()))
        return vacio, vacio.drop(columns=["Fecha"]), vacio.drop(columns=["Fecha", "Semana"])
    e["SEM"] = e["FECHA"].map(inicio_semana)
    etiquetas = {sem: etiqueta_semana(list(g["FECHA"].unique())) for sem, g in e.groupby("SEM")}
    dias = []
    for fecha, g in e.groupby("FECHA"):
        dias.append({"Fecha": fecha.strftime("%d/%m/%Y"), "Semana": etiquetas[inicio_semana(fecha)], **_fila_acido(g)})
    semanas = [{"Semana": etiquetas[sem], **_fila_acido(g)} for sem, g in e.groupby("SEM")]
    mes = [{**_fila_acido(e)}]
    return pd.DataFrame(dias), pd.DataFrame(semanas), pd.DataFrame(mes)


def tabla_patentes(b):
    """Nivel de servicio de cada patente en el mes: órdenes totales, de Easy y de París, entregadas o recogidas,
    % de entrega del mes (entregadas / totales) y el % de cada cliente."""
    e = b[(b["ES_EASY"] | b["ES_PARIS"]) & b["PAT_KEY"].ne("")]
    cols = ["Patente", "Órdenes totales", "Easy", "París", "Entregadas / recogidas", "No entregadas",
            "% entrega del mes", "NS Easy (%)", "NS París (%)", "Días en ruta"]
    if e.empty:
        return pd.DataFrame(columns=cols)
    e = e.assign(EASY_ENT=e["ES_EASY"] & e["ENTREGADO"], PARIS_ENT=e["ES_PARIS"] & e["ENTREGADO"])
    g = e.groupby("PAT_KEY").agg(Patente=("PATENTE", "first"), tot=("ORDEN", "size"), easy=("ES_EASY", "sum"),
                                 paris=("ES_PARIS", "sum"), ent=("ENTREGADO", "sum"), easy_ent=("EASY_ENT", "sum"),
                                 paris_ent=("PARIS_ENT", "sum"), dias=("FECHA", "nunique")).reset_index(drop=True)
    out = pd.DataFrame({
        "Patente": g["Patente"], "Órdenes totales": g["tot"].astype(int), "Easy": g["easy"].astype(int),
        "París": g["paris"].astype(int), "Entregadas / recogidas": g["ent"].astype(int),
        "No entregadas": (g["tot"] - g["ent"]).astype(int),
        "% entrega del mes": (g["ent"] / g["tot"] * 100).round(2),
        "NS Easy (%)": (g["easy_ent"] / g["easy"].where(g["easy"] > 0) * 100).round(2),
        "NS París (%)": (g["paris_ent"] / g["paris"].where(g["paris"] > 0) * 100).round(2),
        "Días en ruta": g["dias"].astype(int),
    })
    return out.sort_values("Órdenes totales", ascending=False).reset_index(drop=True)


def tabla_subestados(b, solo_no_entregados=False):
    """Cantidad de órdenes de cada sub-estado en el mes (solo Easy y París)."""
    base = b[b["ES_EASY"] | b["ES_PARIS"]]
    e = base[~base["ENTREGADO"]] if solo_no_entregados else base
    nombre_pct = "% de las no entregadas" if solo_no_entregados else "% de las órdenes del mes"
    cols = ["Estado", "Sub-estado", "Cantidad", nombre_pct]
    if e.empty:
        return pd.DataFrame(columns=cols)
    e = e.assign(SUB=e["SUBESTADO"].replace("", "(sin sub-estado)"),
                 SUBK=e["SUB_CLEAN"].replace("", "(SIN SUB-ESTADO)"))
    g = (e.groupby(["ESTADO", "SUBK"]).agg(Estado=("ESTADO_TXT", "first"), Sub=("SUB", "first"),
                                          n=("ORDEN", "size")).reset_index(drop=True))
    out = pd.DataFrame({"Estado": g["Estado"], "Sub-estado": g["Sub"], "Cantidad": g["n"].astype(int),
                        nombre_pct: (g["n"] / len(e) * 100).round(2)})
    return out.sort_values("Cantidad", ascending=False).reset_index(drop=True)


PARIS_MENSUAL_CAMPOS = {
    **PARIS_CAMPOS,
    "FECHA": {"etiqueta": "Fecha del día", "exactos": ["Fecha"], "pos": 0, "req": True},
}


def procesar_calce_mensual(y, p_raw, anio, mes):
    """Calce París de todo el mes. Por cada orden se toma su último estado en cada sistema (el más reciente del
    mes) y se compara. Las que siguen sin coincidir son las órdenes que aún no se arreglan."""
    mapa_p, avisos_p = mapear_columnas(p_raw, PARIS_MENSUAL_CAMPOS)
    en_mes = lambda serie: serie.map(lambda d: (not pd.isna(d)) and d.year == anio and d.month == mes)  # noqa: E731
    p = pd.DataFrame({
        "ORDEN": txt(_col(p_raw, mapa_p, "ORDEN")), "EST_P": txt(_col(p_raw, mapa_p, "ESTADO")),
        "SUB_P": txt(_col(p_raw, mapa_p, "SUBESTADO")),
        "FECHA_P": pd.to_datetime(_col(p_raw, mapa_p, "FECHA"), errors="coerce", dayfirst=True,
                                  format="mixed").dt.date,
    })
    p = p[en_mes(p["FECHA_P"])]
    yp = y[y["ES_PARIS"] & en_mes(y["FECHA"])]
    if yp.empty:
        raise ValueError(f"El BTK de Yáñez no tiene pedidos de París en {MESES[mes - 1]} de {anio}.")
    if p.empty:
        raise ValueError(f"El BTK de París no tiene filas de {MESES[mes - 1]} de {anio}.")

    yp = yp.sort_values("FECHA", kind="stable")
    hist = yp.groupby("ORDEN")["FECHA"].agg(PRIMER="min", ULTIMO="max")
    y_last = yp.drop_duplicates("ORDEN", keep="last")
    p = p.sort_values("FECHA_P", kind="stable")
    p_last = p.drop_duplicates("ORDEN", keep="last")
    m = y_last.merge(p_last, on="ORDEN", how="left").merge(hist, left_on="ORDEN", right_index=True, how="left")
    m["ENCONTRADO"] = m["EST_P"].notna()
    m["EST_P_C"] = m["EST_P"].fillna("").map(limpiar_texto)
    m["SUB_P_C"] = m["SUB_P"].fillna("").map(limpiar_texto)
    res = [evaluar_calce(e, s_, ep, sp, enc) for e, s_, ep, sp, enc in
           zip(m["ESTADO"], m["SUB_CLEAN"], m["EST_P_C"], m["SUB_P_C"], m["ENCONTRADO"])]
    m["OK"] = [r[0] for r in res]
    m["OBS"] = [r[1] for r in res]
    fecha_ref = yp["FECHA"].max()

    todas = pd.DataFrame({
        "N° de pedido": m["ORDEN"],
        "Fecha Yáñez": m["FECHA"].map(fmt_fecha),
        "Beetrack T.Yañez Estado": m["ESTADO_TXT"],
        "Beetrack T.Yañez Sub Estado": m["SUBESTADO"],
        "Beetrack París": m["EST_P"].fillna("NO ENCONTRADO"),
        "Beetrack París Sub Estado": m["SUB_P"].fillna("NO ENCONTRADO"),
        "Fecha París": m["FECHA_P"].map(fmt_fecha),
        "Estado Cruce": m["OK"].map({True: "Coincide", False: "No Coincide"}),
        "Observación": "",
        "Fecha Compromiso": m["F_COMP"].map(fmt_fecha),
        "Diagnóstico automático": m["OBS"],
        "Primer día en Yáñez": m["PRIMER"].map(fmt_fecha),
        "Días desde su primer día": m["PRIMER"].map(lambda d: (fecha_ref - d).days),
    })
    orden_f = m["FECHA"].values
    todas = todas.iloc[pd.Series(orden_f).argsort(kind="stable").values].reset_index(drop=True)
    pendientes = todas[todas["Estado Cruce"] == "No Coincide"].sort_values(
        "Días desde su primer día", ascending=False, kind="stable").reset_index(drop=True)

    dia = (todas.assign(_f=m["PRIMER"].values).groupby("_f")["Estado Cruce"]
           .agg(Órdenes="size", Pendientes=lambda s_: int((s_ == "No Coincide").sum())).reset_index())
    dia["_f"] = dia["_f"].map(fmt_fecha)
    dia = dia.rename(columns={"_f": "Primer día en Yáñez"})
    dia["% coincidencia"] = ((dia["Órdenes"] - dia["Pendientes"]) / dia["Órdenes"] * 100).round(2)
    dia = dia.rename(columns={"Órdenes": "Órdenes cruzadas"})
    diag = (pendientes.groupby("Diagnóstico automático").size().reset_index(name="Cantidad")
            .sort_values("Cantidad", ascending=False).reset_index(drop=True))
    comb = (pendientes.groupby(["Beetrack T.Yañez Estado", "Beetrack T.Yañez Sub Estado", "Beetrack París",
                                "Beetrack París Sub Estado"]).size().reset_index(name="Cantidad")
            .sort_values("Cantidad", ascending=False).reset_index(drop=True))
    total, pend = int(len(todas)), int(len(pendientes))
    meta = {"anio": int(anio), "mes": int(mes), "ordenes": total, "pendientes": pend, "coinciden": total - pend,
            "pct_coincidencia": round(pct(total - pend, total), 2),
            "no_encontradas": int((todas["Beetrack París"] == "NO ENCONTRADO").sum()),
            "solo_paris": int(len(set(p_last["ORDEN"]) - set(y_last["ORDEN"]))),
            "mapeo_paris": mapa_a_filas(mapa_p), "avisos_paris": list(avisos_p),
            "mapeo_btk": mapa_a_filas(y.attrs.get("mapeo", {})), "avisos_btk": list(y.attrs.get("avisos", []))}
    return meta, {"calce_pendientes": pendientes, "calce_todas": todas, "calce_dia": dia,
                  "calce_diagnostico": diag, "calce_combinaciones": comb}


def _nombre_conductor(etiqueta):
    """«RUTA A - Raul» -> «Raul»."""
    return re.sub(r"(?i)^\s*ruta\s+[A-Z]{1,2}\s*-\s*", "", str(etiqueta)).strip()


def nombres_habituales(b):
    """Para cada patente, los nombres de driver que usa en el período:
    {clave de patente: [(clave del nombre, nombre, proporción de sus pedidos, días en que fue el nombre principal)]}
    de mayor a menor proporción."""
    e = b[(b["ES_EASY"] | b["ES_PARIS"]) & b["PAT_KEY"].ne("") & b["CONDUCTOR"].ne("")].copy()
    e["NOMBRE"] = e["CONDUCTOR"].map(_nombre_conductor)
    e = e[e["NOMBRE"].ne("") & ~e["NOMBRE"].str.upper().str.contains("REINGRESO")]
    e["CLAVE"] = e["NOMBRE"].map(limpiar_texto)
    # nombre principal de cada patente en cada día
    por_dia = e.groupby(["PAT_KEY", "FECHA", "CLAVE"]).size().reset_index(name="n")
    principal = por_dia.sort_values("n", ascending=False).drop_duplicates(["PAT_KEY", "FECHA"])
    dias_principal = principal.groupby(["PAT_KEY", "CLAVE"]).size()
    out = {}
    for pk, g_ in e.groupby("PAT_KEY"):
        cuenta = g_["CLAVE"].value_counts()
        out[pk] = [(clave, g_.loc[g_["CLAVE"] == clave, "NOMBRE"].value_counts().index[0].title(), n / cuenta.sum(),
                    int(dias_principal.get((pk, clave), 0))) for clave, n in cuenta.items()]
    return out


def nombre_del_dia(tabla_pat, bp, min_dias=2):
    """Driver de una patente en un día. Entre los nombres que esa patente maneja de verdad (fueron el nombre
    principal al menos `min_dias` días del mes), se elige el que más aparece ese día. Si ninguno aparece, se usa el
    más habitual del mes. Así una etiqueta suelta de otro driver (pedidos que cambiaron de ruta) no define quién
    manejó."""
    if not tabla_pat:
        return None
    validos = {c: nombre for c, nombre, _, dias in tabla_pat if dias >= min_dias}
    dia = bp["CONDUCTOR"][bp["CONDUCTOR"].ne("")].map(_nombre_conductor).map(limpiar_texto).value_counts()
    for clave in dia.index:
        if clave in validos:
            return validos[clave]
    return tabla_pat[0][1]


def comunas_mensual(hela, b, anio, mes):
    """Reporte de comunas de todo el mes: aplica el cruce Hela + BTK día por día."""
    h = preparar_hela(hela)
    if "FECHA" not in h.attrs["mapeo"]:
        raise ValueError("No se encontró la columna de fecha en el Hela mensual (se espera «Fecha» en la columna A).")
    f_hela = pd.to_datetime(h["FECHA"], errors="coerce", dayfirst=True, format="mixed").ffill().dt.date
    dias_hela = {d for d in f_hela.dropna().unique() if d.year == anio and d.month == mes}
    dias_btk = {d for d in b["FECHA"].dropna().unique() if d.year == anio and d.month == mes}
    avisos, piezas = [], []
    sin_hela = sorted(dias_btk - dias_hela)
    sin_btk = sorted(dias_hela - dias_btk)
    if sin_hela:
        avisos.append("Días con BTK y sin Hela (sus patentes salen como «Solo BTK»): "
                      + ", ".join(f"{d:%d/%m}" for d in sin_hela))
    if sin_btk:
        avisos.append("Días con Hela y sin pedidos en el BTK (sus rutas salen como «Solo Hela»): "
                      + ", ".join(f"{d:%d/%m}" for d in sin_btk))
    nombres = nombres_habituales(b)
    for dia in sorted(dias_hela | dias_btk):
        try:
            df_dia, info_dia = procesar_comunas(hela[(f_hela == dia).values], b[b["FECHA"] == dia], dia, nombres)
            piezas.append(df_dia)
            avisos.extend(f"{dia:%d/%m}: {a}" for a in info_dia.get("avisos_rutas", []))
        except Exception as e:  # noqa: BLE001
            avisos.append(f"{dia:%d/%m}: no se pudo procesar ({e})")
    df = pd.concat(piezas, ignore_index=True) if piezas else pd.DataFrame()
    return df, avisos, len(dias_hela | dias_btk)


def procesar_mensual(hela, b, anio, mes):
    """Procesa el mes completo. Devuelve (meta, dfs) listos para guardar."""
    if b["FECHA"].isna().all():
        raise ValueError("No se encontró la columna de fecha en el BTK mensual (se espera «Fecha» en la columna A).")
    en_mes = b["FECHA"].map(lambda d: (not pd.isna(d)) and d.year == anio and d.month == mes)
    fuera = int((b["FECHA"].notna() & ~en_mes).sum())
    bm = b[en_mes]
    if bm.empty:
        raise ValueError(f"El BTK no tiene pedidos de {MESES[mes - 1]} de {anio}.")
    comunas, avisos, n_dias = comunas_mensual(hela, bm, anio, mes)
    if fuera:
        avisos.append(f"El BTK trae {fuera} fila(s) de otros meses; no se consideraron.")
    dia, semana, mes_df = tablas_ns_acido(bm)
    meta = {"anio": int(anio), "mes": int(mes), "dias": int(n_dias), "avisos": avisos,
            "mapeo_btk": mapa_a_filas(b.attrs.get("mapeo", {})), "avisos_btk": list(b.attrs.get("avisos", []))}
    dfs = {"comunas": comunas, "acido_dia": dia, "acido_semana": semana, "acido_mes": mes_df,
           "patentes": tabla_patentes(bm), "subestados_noent": tabla_subestados(bm, True)}
    return meta, dfs


# --- NS compromiso: datos que se ingresan a mano por día ---
# Fecha compromiso = pedidos cuya fecha de entrega límite es ese mismo día (aunque se entreguen después, incluso
# si el pedido venía de un mes anterior). Por eso no se calcula desde el BTK: se ingresa por cliente.
def guardar_fc(fecha, easy, paris):
    """easy y paris son (totales, entregados, no entregados). El general es la suma."""
    t, e, n = (easy[i] + paris[i] for i in range(3))
    meta = {"fc_tot": int(t), "fc_ent": int(e), "fc_no": int(n)}
    for pref, (a, b_, c_) in (("fc_easy", easy), ("fc_paris", paris)):
        meta[f"{pref}_tot"], meta[f"{pref}_ent"], meta[f"{pref}_no"] = int(a), int(b_), int(c_)
    guardar_resultado("fc_manual", fecha, meta, {})


def cargar_fc(f_ini, f_fin):
    """Lista de dicts {fecha, tot, ent, no, easy, paris}; easy/paris son (tot, ent, no) o None si el día
    solo tiene un total general guardado."""
    out = []
    for f, m, _ in cargar("fc_manual", f_ini, f_fin):
        def par(pref):
            if f"{pref}_tot" in m:
                return int(m[f"{pref}_tot"]), int(m[f"{pref}_ent"]), int(m[f"{pref}_no"])
            return None

        out.append({"fecha": f, "tot": int(m.get("fc_tot", 0)), "ent": int(m.get("fc_ent", 0)),
                    "no": int(m.get("fc_no", 0)), "easy": par("fc_easy"), "paris": par("fc_paris"),
                    "origen": "Mensual"})
    return out


def cargar_fc_mes(f_ini, f_fin):
    """Datos de compromiso del rango. Lo ingresado en la pestaña mensual manda; los días que no tienen nada se
    completan con lo que ya ingresaste en el NS diario (pestaña 1), para no tener que ingresarlo dos veces."""
    regs = cargar_fc(f_ini, f_fin)
    tienen = {r["fecha"] for r in regs}
    for f, m, _ in cargar("ns", f_ini, f_fin):
        if f in tienen:
            continue
        c = calcular_ns(m)
        if c["fc_gen_tot"]:
            regs.append({"fecha": f, "tot": c["fc_gen_tot"], "ent": c["fc_gen_ent"], "no": c["fc_gen_no"],
                         "easy": (c["fc_easy_tot"], c["fc_easy_ent"], c["fc_easy_no"]),
                         "paris": (c["fc_paris_tot"], c["fc_paris_ent"], c["fc_paris_no"]),
                         "origen": "NS diario"})
    return sorted(regs, key=lambda r: r["fecha"])


def borrar_fc(fecha):
    STORE.borrar_resultado("fc_manual", fecha.isoformat())
    _invalidar()


def _ns_o_nan(ent, tot):
    return round(pct(ent, tot), 2) if tot else float("nan")


def tablas_ns_compromiso(regs):
    """Solo porcentajes: NS compromiso = entregados / totales, por día, por semana y del mes.
    Devuelve (por_dia, por_semana, mes) donde mes es {"general", "easy", "paris"} o None si no hay datos."""
    cols_d = ["Fecha", "Semana", "NS Compromiso (%)", "NS Easy (%)", "NS París (%)"]
    cols_s = ["Semana", "NS Compromiso (%)", "NS Easy (%)", "NS París (%)"]
    regs = sorted(regs, key=lambda r: r["fecha"])
    if not regs:
        return pd.DataFrame(columns=cols_d), pd.DataFrame(columns=cols_s), None

    def suma(rs, cli):
        """(entregados, totales) sumados; para Easy/París solo cuentan los días que traen ese cliente."""
        if cli is None:
            return sum(r["ent"] for r in rs), sum(r["tot"] for r in rs)
        con = [r[cli] for r in rs if r[cli] is not None]
        return sum(x[1] for x in con), sum(x[0] for x in con)

    sems = {}
    for r in regs:
        sems.setdefault(inicio_semana(r["fecha"]), []).append(r)
    etiquetas = {k: etiqueta_semana([r["fecha"] for r in v]) for k, v in sems.items()}
    dias = [{"Fecha": r["fecha"].strftime("%d/%m/%Y"), "Semana": etiquetas[inicio_semana(r["fecha"])],
             "NS Compromiso (%)": _ns_o_nan(r["ent"], r["tot"]),
             "NS Easy (%)": _ns_o_nan(*suma([r], "easy")), "NS París (%)": _ns_o_nan(*suma([r], "paris"))}
            for r in regs]
    semanas = [{"Semana": etiquetas[k], "NS Compromiso (%)": _ns_o_nan(*suma(v, None)),
                "NS Easy (%)": _ns_o_nan(*suma(v, "easy")), "NS París (%)": _ns_o_nan(*suma(v, "paris"))}
               for k, v in sorted(sems.items())]
    mes = {"general": _ns_o_nan(*suma(regs, None)), "easy": _ns_o_nan(*suma(regs, "easy")),
           "paris": _ns_o_nan(*suma(regs, "paris"))}
    return pd.DataFrame(dias, columns=cols_d), pd.DataFrame(semanas, columns=cols_s), mes


def excel_mensual(dfs, regs, anio, mes):
    hojas = {}
    com = dfs.get("comunas", pd.DataFrame())
    if not com.empty:
        hojas["Comunas del mes"] = com
        bon = df_bonificacion(com)
        hojas["Bonificación del mes"] = bon
        hojas["Bonificación por patente"] = bonificacion_por_patente(bon)
    hojas.update(hojas_compromiso(regs, anio, mes))
    for clave, nombre in (("acido_dia", "NS Ácido día"), ("acido_semana", "NS Ácido semana"),
                          ("acido_mes", "NS Ácido mes"), ("patentes", "Órdenes y NS por patente"),
                          ("subestados_noent", "Sub-estados no entregados"),
                          ("calce_pendientes", "Calce París pendientes"), ("calce_todas", "Calce París todas"),
                          ("calce_dia", "Calce París por día"), ("calce_diagnostico", "Calce París diagnóstico")):
        d = dfs.get(clave)
        if d is not None and not d.empty:
            hojas[nombre] = d
    return a_excel(hojas, {**GRAFICO_ACIDO_EXCEL, **GRAFICO_COMPROMISO_EXCEL}) if hojas else None


# --- PDF de los reportes mensuales ---
def _celda_pdf(col, v):
    try:
        if v is None or pd.isna(v):
            return "—"
    except (TypeError, ValueError):
        pass
    if _es_pct(col):
        try:
            return fmt_pct(float(v))
        except (TypeError, ValueError):
            return str(v)
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, numbers.Integral) or (isinstance(v, numbers.Real) and float(v).is_integer()):
        return f"{int(v):,}".replace(",", ".")
    if isinstance(v, numbers.Real):
        return f"{float(v):.2f}".replace(".", ",")
    return str(v)


def _grafico_lineas_pdf(g, ancho):
    """Gráfico de líneas (una línea por serie) dibujado con reportlab, sin librerías extra.
    g = {"df": DataFrame, "x": columna de fechas, "series": [columnas en %], "colores": [#hex], "etiquetas": [...]}"""
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.widgets.markers import makeMarker
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    d = g["df"]
    cats = [str(x)[:5] for x in d[g["x"]]]
    datos = [[None if pd.isna(v) else float(v) for v in d[c]] for c in g["series"]]
    todos = [v for fila in datos for v in fila if v is not None]
    ymin = max(0, int((min(todos) - 5) // 10 * 10)) if todos else 0
    alto = 6.4 * cm
    dib = Drawing(ancho, alto)
    lc = HorizontalLineChart()
    lc.x, lc.y, lc.width, lc.height = 34, 46, ancho - 48, alto - 56
    lc.data = datos
    lc.joinedLines = 1
    lc.categoryAxis.categoryNames = cats
    lc.categoryAxis.labels.angle = 90
    lc.categoryAxis.labels.boxAnchor = "e"
    lc.categoryAxis.labels.fontSize = 6.5
    lc.categoryAxis.labels.dy = -2
    lc.categoryAxis.strokeColor = colors.HexColor("#9AA7BF")
    lc.valueAxis.valueMin, lc.valueAxis.valueMax = ymin, 100
    lc.valueAxis.valueStep = 5 if (100 - ymin) <= 30 else 10
    lc.valueAxis.labelTextFormat = "%d%%"
    lc.valueAxis.labels.fontSize = 7
    lc.valueAxis.strokeColor = colors.HexColor("#9AA7BF")
    lc.valueAxis.visibleGrid = 1
    lc.valueAxis.gridStrokeColor = colors.HexColor("#E3E8F0")
    for i, col in enumerate(g["colores"]):
        lc.lines[i].strokeColor = colors.HexColor(col)
        lc.lines[i].strokeWidth = 1.6
        lc.lines[i].symbol = makeMarker("Circle")
        lc.lines[i].symbol.size = 3
        lc.lines[i].symbol.fillColor = colors.HexColor(col)
        lc.lines[i].symbol.strokeColor = colors.HexColor(col)
    dib.add(lc)
    leyenda = Legend()
    leyenda.x, leyenda.y = ancho / 2 - 110, 8
    leyenda.alignment = "right"
    leyenda.columnMaximum = 1
    leyenda.deltax, leyenda.deltay = 95, 0
    leyenda.dxTextSpace = 5
    leyenda.fontSize = 8
    leyenda.colorNamePairs = [(colors.HexColor(c), n) for c, n in zip(g["colores"], g["etiquetas"])]
    dib.add(leyenda)
    return dib


def pdf_reporte(titulo, secciones):
    """PDF con tablas. secciones = [{"titulo": str, "metricas": [(etiqueta, valor)], "df": DataFrame, "nota": str}].
    Si alguna tabla es ancha, la hoja sale horizontal."""
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ancho_max = max([len(sec["df"].columns) for sec in secciones if sec.get("df") is not None] or [0])
    pagina = landscape(A4) if ancho_max > 7 else A4
    margen = 1.5 * cm
    disp = pagina[0] - 2 * margen
    buf = io.BytesIO()

    def pie(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6B7A99"))
        canvas.drawString(margen, 0.8 * cm, titulo)
        canvas.drawRightString(pagina[0] - margen, 0.8 * cm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=pagina, leftMargin=margen, rightMargin=margen, topMargin=1.4 * cm,
                            bottomMargin=1.6 * cm, title=titulo, author="Transportes Yáñez")
    est = getSampleStyleSheet()
    s_tit = ParagraphStyle("t", parent=est["Title"], alignment=0, fontSize=16, spaceAfter=2)
    s_sec = ParagraphStyle("s", parent=est["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    s_nota = ParagraphStyle("n", parent=est["Normal"], fontSize=8, textColor=colors.HexColor("#6B7A99"))
    historia = [Paragraph(escape(titulo), s_tit)]

    for sec in secciones:
        bloque = [Paragraph(escape(sec["titulo"]), s_sec)]
        if sec.get("metricas"):
            etiq = [Paragraph(f'<font size="8" color="#6B7A99">{escape(str(e))}</font>', est["Normal"])
                    for e, _ in sec["metricas"]]
            vals = [Paragraph(f'<font size="15"><b>{escape(str(v))}</b></font>', est["Normal"])
                    for _, v in sec["metricas"]]
            m = Table([etiq, vals], colWidths=[min(5.2 * cm, disp / len(etiq))] * len(etiq), hAlign="LEFT")
            m.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C0CC")),
                                   ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE3EC")),
                                   ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F7FB")),
                                   ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            bloque.append(m)
        if sec.get("grafico") is not None:
            bloque.append(_grafico_lineas_pdf(sec["grafico"], disp))
        df = sec.get("df")
        if df is not None and not df.empty:
            cols = list(df.columns)
            fuente = 8 if len(cols) <= 8 else 7
            f_h = ParagraphStyle("h", parent=est["Normal"], fontName="Helvetica-Bold", fontSize=fuente,
                                 leading=fuente + 1.5)
            filas = [[_celda_pdf(c, v) for c, v in zip(cols, fila)] for fila in df.itertuples(index=False, name=None)]
            anch = []
            for i, c in enumerate(cols):
                w_h = max(stringWidth(w_, "Helvetica-Bold", fuente) for w_ in (str(c).split() or [""]))
                w_b = max([stringWidth(f[i], "Helvetica", fuente) for f in filas] or [0])
                anch.append(max(w_h, w_b) + 10)
            if sum(anch) > disp:
                anch = [a * disp / sum(anch) for a in anch]
            datos = [[Paragraph(escape(str(c)), f_h) for c in cols]] + filas
            t = Table(datos, colWidths=anch, repeatRows=1, hAlign="LEFT")
            num = [i for i, c in enumerate(cols) if pd.api.types.is_numeric_dtype(df[c])]
            estilo = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EDF2")),
                      ("FONTNAME", (0, 1), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 1), (-1, -1), fuente),
                      ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C0CC")),
                      ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FC")]),
                      ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                      ("TOPPADDING", (0, 0), (-1, -1), 2.5 if len(df) < 20 else 1.4),
                      ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5 if len(df) < 20 else 1.4),
                      ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
            estilo += [("ALIGN", (i, 1), (i, -1), "RIGHT") for i in num]
            t.setStyle(TableStyle(estilo))
            bloque += [Spacer(1, 4), t]
        if sec.get("nota"):
            bloque += [Spacer(1, 3), Paragraph(escape(sec["nota"]), s_nota)]
        historia.append(KeepTogether(bloque) if (df is None or len(df) < 14) else bloque[0])
        if df is not None and len(df) >= 14:
            historia += bloque[1:]
    doc.build(historia, onFirstPage=pie, onLaterPages=pie)
    return buf.getvalue()


def _nombre_mes(anio, mes):
    return f"{MESES[mes - 1].capitalize()} {anio}"


def GRAFICO_COMPROMISO_PDF(dia_c):  # noqa: N802
    return {"df": dia_c, "x": "Fecha", "series": ["NS Easy (%)", "NS París (%)", "NS Compromiso (%)"],
            "etiquetas": ["NS Easy (%)", "NS París (%)", "NS Compromiso general (%)"],
            "colores": ["#1F5FBF", "#7FC4FF", "#FF2B2B"]}


def _secciones_compromiso(regs, anio, mes):
    dia_c, sem_c, mes_c = tablas_ns_compromiso(regs)
    if mes_c is None:
        return []
    f = lambda x: fmt_pct(x) if x == x else "—"  # noqa: E731
    return [
        {"titulo": "NS Compromiso del mes",
         "metricas": [("General", f(mes_c["general"])), ("Easy", f(mes_c["easy"])), ("París", f(mes_c["paris"]))]},
        {"titulo": "Evolución diaria", "grafico": GRAFICO_COMPROMISO_PDF(dia_c)},
        {"titulo": "Por semana", "df": sem_c},
        {"titulo": "Por día", "df": dia_c},
    ]


def _secciones_acido(dfs):
    if "acido_dia" not in dfs or dfs["acido_dia"].empty:
        return []
    m = dfs["acido_mes"].iloc[0]
    return [
        {"titulo": "NS Ácido del mes",
         "metricas": [("NS Ácido", fmt_pct(m["NS Ácido (%)"])), ("Órdenes totales", _celda_pdf("x", m["Órdenes totales"])),
                      ("Entregadas / recogidas", _celda_pdf("x", m["Entregadas / recogidas"])),
                      ("No entregadas", _celda_pdf("x", m["No entregadas"]))],
         "nota": "NS ácido = órdenes entregadas o recogidas / órdenes totales (todas las órdenes de Easy y París)."},
        {"titulo": "Evolución diaria",
         "grafico": {"df": dfs["acido_dia"], "x": "Fecha", "series": ["NS Easy (%)", "NS París (%)", "NS Ácido (%)"],
                     "etiquetas": ["NS Easy (%)", "NS París (%)", "NS Ácido (%)"],
                     "colores": ["#1F5FBF", "#7FC4FF", "#FF2B2B"]}},
        {"titulo": "Por semana", "df": dfs["acido_semana"]},
        {"titulo": "Por día", "df": dfs["acido_dia"]},
    ]


def pdf_ns_compromiso(regs, anio, mes):
    return pdf_reporte(f"NS Compromiso | {_nombre_mes(anio, mes)}", _secciones_compromiso(regs, anio, mes))


def _secciones_extra_acido(dfs):
    """Lo que acompaña al NS ácido: el nivel de servicio de cada patente y los sub-estados de no entrega."""
    secs = []
    if dfs.get("patentes") is not None and not dfs["patentes"].empty:
        secs.append({"titulo": "NS de cada patente en el mes", "df": dfs["patentes"],
                     "nota": "% entrega del mes = órdenes entregadas o recogidas / órdenes totales de la patente."})
    if dfs.get("subestados_noent") is not None and not dfs["subestados_noent"].empty:
        secs.append({"titulo": "Sub-estados de las órdenes no entregadas", "df": dfs["subestados_noent"]})
    return secs


def pdf_ns_acido(dfs, anio, mes):
    return pdf_reporte(f"NS Ácido | {_nombre_mes(anio, mes)}", _secciones_acido(dfs) + _secciones_extra_acido(dfs))


def hojas_acido(dfs):
    """Hojas de Excel del NS ácido: por día, por semana, del mes, por patente y sub-estados de no entrega."""
    hojas = {"NS Ácido día": dfs["acido_dia"], "NS Ácido semana": dfs["acido_semana"], "NS Ácido mes": dfs["acido_mes"]}
    if dfs.get("patentes") is not None and not dfs["patentes"].empty:
        hojas["NS por patente del mes"] = dfs["patentes"]
    if dfs.get("subestados_noent") is not None and not dfs["subestados_noent"].empty:
        hojas["Sub-estados no entregados"] = dfs["subestados_noent"]
    return hojas


def pdf_patentes(dfs, anio, mes):
    return pdf_reporte(f"NS por patente | {_nombre_mes(anio, mes)}", [
        {"titulo": "Nivel de servicio de cada patente en el mes", "df": dfs["patentes"],
         "nota": "% entrega del mes = órdenes entregadas o recogidas / órdenes totales de la patente (solo Easy y París)."}])


def pdf_subestados(dfs, anio, mes):
    se = dfs["subestados_noent"]
    return pdf_reporte(f"Sub-estados de no entrega | {_nombre_mes(anio, mes)}", [
        {"titulo": "Órdenes no entregadas, por sub-estado",
         "metricas": [("Total no entregadas", _celda_pdf("x", int(se["Cantidad"].sum())))], "df": se}])


def pdf_bonificacion(com, anio, mes):
    bon = df_bonificacion(com)
    por_pat = bonificacion_por_patente(bon)
    return pdf_reporte(f"Bonificación | {_nombre_mes(anio, mes)}", [
        {"titulo": "Resumen del mes",
         "metricas": [("Total bonificación", "$ " + _celda_pdf("x", int(bon["D: Valor"].sum()))),
                      ("Rutas con pago", _celda_pdf("x", len(bon))), ("Patentes", _celda_pdf("x", len(por_pat)))]},
        {"titulo": "Total por patente", "df": por_pat}])


def pdf_mensual_completo(dfs, regs, anio, mes):
    """Un solo PDF con NS compromiso, NS ácido, NS por patente, sub-estados de no entrega y bonificación."""
    secs = _secciones_compromiso(regs, anio, mes) + _secciones_acido(dfs) + _secciones_extra_acido(dfs)
    if dfs.get("comunas") is not None and not dfs["comunas"].empty:
        bon = df_bonificacion(dfs["comunas"])
        secs.append({"titulo": "Bonificación del mes por patente",
                     "metricas": [("Total bonificación", "$ " + _celda_pdf("x", int(bon["D: Valor"].sum()))),
                                  ("Rutas con pago", _celda_pdf("x", len(bon)))],
                     "df": bonificacion_por_patente(bon)})
    return pdf_reporte(f"Reporte mensual | {_nombre_mes(anio, mes)}", secs)


def hojas_compromiso(regs, anio, mes):
    """Hojas de Excel del NS compromiso: por día, por semana, del mes y los datos ingresados."""
    dia_c, sem_c, mes_c = tablas_ns_compromiso(regs)
    if mes_c is None:
        return {}
    hojas = {"NS Compromiso día": dia_c, "NS Compromiso semana": sem_c,
             "NS Compromiso mes": pd.DataFrame([{"Mes": _nombre_mes(anio, mes), "NS Compromiso (%)": mes_c["general"],
                                                "NS Easy (%)": mes_c["easy"], "NS París (%)": mes_c["paris"]}])}
    filas = []
    for r in sorted(regs, key=lambda r: r["fecha"]):
        fila = {"Fecha": r["fecha"].strftime("%d/%m/%Y"), "Origen": r.get("origen", "Mensual")}
        for nombre, clave in (("Easy", "easy"), ("París", "paris")):
            t, e, n = r[clave] if r[clave] is not None else (None, None, None)
            fila.update({f"{nombre} totales": t, f"{nombre} entregados": e, f"{nombre} no entregados": n})
        fila.update({"General totales": r["tot"], "General entregados": r["ent"], "General no entregados": r["no"]})
        filas.append(fila)
    hojas["Datos compromiso ingresados"] = pd.DataFrame(filas)
    return hojas


# ===================== INTERFAZ =====================
# =====================================================================
# ACCESO CON CONTRASEÑA
# =====================================================================
# =====================================================================
# PANTALLA DE INICIO DE SESIÓN (tema espacial con gato y perro astronauta)
# =====================================================================
def _estrellas(semilla, n, tam_max):
    """Campo de estrellas como lista de box-shadow (posiciones en vw/vh)."""
    r = random.Random(semilla)
    sombras = []
    for _ in range(n):
        x, y = r.uniform(0, 100), r.uniform(0, 92)
        op = r.uniform(0.35, 1)
        tono = r.choice(["255,255,255", "255,255,255", "200,225,255", "170,200,255", "255,240,220"])
        sombras.append(f"{x:.1f}vw {y:.1f}vh 0 {r.choice([0, 0, 0.5, 1]) * tam_max / 2:.1f}px rgba({tono},{op:.2f})")
    return ",".join(sombras)


def _astronauta(uid, cabeza, acento, suela):
    """Astronauta con traje blanco; `cabeza` es el dibujo de la cara dentro del casco."""
    return f"""
<svg viewBox="0 0 220 300" xmlns="http://www.w3.org/2000/svg">
<defs>
<linearGradient id="traje{uid}" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#ffffff"/><stop offset="1" stop-color="#cfd9ec"/></linearGradient>
<linearGradient id="mochila{uid}" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#b9c4da"/><stop offset="1" stop-color="#7f8fb0"/></linearGradient>
<radialGradient id="vidrio{uid}" cx="0.32" cy="0.26" r="0.95"><stop offset="0" stop-color="#ffffff" stop-opacity="0.55"/><stop offset="0.45" stop-color="#8cc4ff" stop-opacity="0.16"/><stop offset="1" stop-color="#1c3f8f" stop-opacity="0.32"/></radialGradient>
<clipPath id="casco{uid}"><circle cx="110" cy="86" r="56"/></clipPath>
</defs>
<rect x="46" y="118" width="128" height="98" rx="26" fill="url(#mochila{uid})"/>
<rect x="60" y="205" width="30" height="8" rx="4" fill="#64729a"/><rect x="130" y="205" width="30" height="8" rx="4" fill="#64729a"/>
<rect x="76" y="196" width="32" height="60" rx="15" fill="url(#traje{uid})"/><rect x="112" y="196" width="32" height="60" rx="15" fill="url(#traje{uid})"/>
<rect x="70" y="246" width="42" height="26" rx="12" fill="{suela}"/><rect x="108" y="246" width="42" height="26" rx="12" fill="{suela}"/>
<rect x="70" y="246" width="42" height="9" rx="4" fill="{acento}"/><rect x="108" y="246" width="42" height="9" rx="4" fill="{acento}"/>
<rect x="64" y="124" width="92" height="98" rx="36" fill="url(#traje{uid})" stroke="#b8c5df" stroke-width="2"/>
<rect x="86" y="150" width="48" height="36" rx="9" fill="#dfe7f5" stroke="#aab9d8" stroke-width="2"/>
<circle cx="97" cy="162" r="4.5" fill="#ff5d5d"/><circle cx="110" cy="162" r="4.5" fill="#46d39a"/><circle cx="123" cy="162" r="4.5" fill="#4aa3ff"/>
<rect x="94" y="173" width="32" height="5" rx="2.5" fill="{acento}"/>
<g transform="rotate(32 62 140)"><rect x="24" y="128" width="50" height="26" rx="13" fill="url(#traje{uid})" stroke="#b8c5df" stroke-width="2"/><circle cx="26" cy="141" r="15" fill="{acento}"/></g>
<g transform="rotate(-62 158 138)"><rect x="146" y="124" width="54" height="26" rx="13" fill="url(#traje{uid})" stroke="#b8c5df" stroke-width="2"/><circle cx="200" cy="137" r="15" fill="{acento}"/></g>
<ellipse cx="110" cy="132" rx="50" ry="13" fill="#c4cfe6" stroke="#a3b2d4" stroke-width="2"/>
<circle cx="110" cy="86" r="56" fill="#16264f"/>
<g clip-path="url(#casco{uid})">{cabeza}</g>
<circle cx="110" cy="86" r="56" fill="url(#vidrio{uid})"/>
<path d="M70 64 A46 46 0 0 1 104 42" fill="none" stroke="#ffffff" stroke-opacity="0.75" stroke-width="6" stroke-linecap="round"/>
<circle cx="76" cy="78" r="4" fill="#ffffff" fill-opacity="0.6"/>
<circle cx="110" cy="86" r="60" fill="none" stroke="#e9f0ff" stroke-width="8"/>
<circle cx="110" cy="86" r="64.5" fill="none" stroke="#9fb0d6" stroke-width="2"/>
<rect x="100" y="14" width="20" height="10" rx="5" fill="#9fb0d6"/>
</svg>"""


CABEZA_GATO = """
<rect x="0" y="0" width="220" height="200" fill="#0d1b3f"/>
<circle cx="170" cy="40" r="1.6" fill="#fff"/><circle cx="60" cy="126" r="1.3" fill="#fff"/><circle cx="150" cy="128" r="1.1" fill="#fff"/>
<polygon points="70,80 66,38 100,58" fill="#f39a3b"/><polygon points="150,80 154,38 120,58" fill="#f39a3b"/>
<polygon points="74,72 72,47 92,60" fill="#ffb7c5"/><polygon points="146,72 148,47 128,60" fill="#ffb7c5"/>
<ellipse cx="110" cy="92" rx="46" ry="40" fill="#f7a94a"/>
<path d="M110 54 L110 68 M96 57 L99 69 M124 57 L121 69" stroke="#d9791f" stroke-width="4" stroke-linecap="round"/>
<path d="M66 90 l12 2 M66 101 l12 -2 M154 90 l-12 2 M154 101 l-12 -2" stroke="#d9791f" stroke-width="4" stroke-linecap="round"/>
<ellipse cx="110" cy="108" rx="24" ry="17" fill="#fff3e2"/>
<ellipse cx="90" cy="88" rx="9" ry="11" fill="#1b1b2f"/><ellipse cx="130" cy="88" rx="9" ry="11" fill="#1b1b2f"/>
<circle cx="93" cy="84" r="3.6" fill="#fff"/><circle cx="133" cy="84" r="3.6" fill="#fff"/><circle cx="88" cy="92" r="1.8" fill="#fff"/><circle cx="128" cy="92" r="1.8" fill="#fff"/>
<polygon points="104,98 116,98 110,106" fill="#ff8fa3"/>
<path d="M110 106 L110 111 M110 111 Q103 119 96 112 M110 111 Q117 119 124 112" fill="none" stroke="#7a3b16" stroke-width="2.6" stroke-linecap="round"/>
<path d="M74 102 l-22 -5 M74 108 l-22 4 M146 102 l22 -5 M146 108 l22 4" stroke="#ffffff" stroke-opacity="0.9" stroke-width="1.8" stroke-linecap="round"/>
<circle cx="76" cy="108" r="6" fill="#ff9aa8" fill-opacity="0.45"/><circle cx="144" cy="108" r="6" fill="#ff9aa8" fill-opacity="0.45"/>
"""

CABEZA_PERRO = """
<rect x="0" y="0" width="220" height="200" fill="#0d1b3f"/>
<circle cx="168" cy="44" r="1.6" fill="#fff"/><circle cx="58" cy="130" r="1.3" fill="#fff"/><circle cx="152" cy="126" r="1.1" fill="#fff"/>
<ellipse cx="66" cy="96" rx="15" ry="31" fill="#8a5727" transform="rotate(14 66 96)"/>
<ellipse cx="154" cy="96" rx="15" ry="31" fill="#8a5727" transform="rotate(-14 154 96)"/>
<ellipse cx="110" cy="90" rx="42" ry="40" fill="#e0a864"/>
<ellipse cx="127" cy="80" rx="19" ry="18" fill="#9b6330" fill-opacity="0.85"/>
<ellipse cx="110" cy="110" rx="26" ry="20" fill="#fff1dc"/>
<ellipse cx="92" cy="84" rx="8" ry="10" fill="#1b1b2f"/><ellipse cx="128" cy="84" rx="8" ry="10" fill="#1b1b2f"/>
<circle cx="95" cy="80" r="3.4" fill="#fff"/><circle cx="131" cy="80" r="3.4" fill="#fff"/>
<path d="M82 70 Q92 64 102 69 M118 69 Q128 64 138 70" fill="none" stroke="#7a4a1f" stroke-width="3" stroke-linecap="round"/>
<ellipse cx="110" cy="100" rx="11" ry="7.5" fill="#2a2030"/><ellipse cx="106" cy="98" rx="3.4" ry="1.8" fill="#fff" fill-opacity="0.6"/>
<path d="M110 107 L110 113 M110 113 Q101 122 94 114 M110 113 Q119 122 126 114" fill="none" stroke="#5a3318" stroke-width="2.6" stroke-linecap="round"/>
<path d="M102 119 Q110 136 118 119 Z" fill="#ff7f93"/>
<path d="M110 119 L110 128" stroke="#e0566f" stroke-width="2"/>
"""


def login_html():
    estrellas_a = _estrellas(11, 150, 2)
    estrellas_b = _estrellas(23, 45, 3)
    gato = _astronauta("g", CABEZA_GATO, "#ff8a3d", "#4a5a86")
    perro = _astronauta("p", CABEZA_PERRO, "#35c0ff", "#4a5a86")
    css = """
<style>
@keyframes flotar{0%,100%{transform:translateY(0) rotate(-4deg)}50%{transform:translateY(-16px) rotate(4deg)}}
@keyframes flotar2{0%,100%{transform:translateY(-10px) rotate(5deg)}50%{transform:translateY(8px) rotate(-3deg)}}
@keyframes titilar{0%,100%{opacity:.35}50%{opacity:1}}
@keyframes pulso{0%,100%{transform:translate(-50%,-50%) scale(1);opacity:.95}50%{transform:translate(-50%,-50%) scale(1.07);opacity:1}}
.stApp{background:radial-gradient(ellipse 120% 70% at 54% 4%,#0c2a66 0%,#06112e 45%,#02050f 100%) !important;color:#dbe8ff}
[data-testid="stAppViewContainer"],[data-testid="stMain"],.main{background:transparent !important}
[data-testid="stHeader"]{background:transparent !important}
[data-testid="stSidebar"],footer,#MainMenu{display:none !important}
.stApp::before{content:"";position:fixed;left:0;top:0;width:1px;height:1px;background:transparent;z-index:0;box-shadow:ESTRELLAS_A}
.stApp::after{content:"";position:fixed;left:0;top:0;width:2px;height:2px;border-radius:50%;background:transparent;z-index:0;box-shadow:ESTRELLAS_B;animation:titilar 3.2s ease-in-out infinite}
.estrella-mayor{position:fixed;left:54%;top:7%;width:min(46vw,520px);aspect-ratio:1;transform:translate(-50%,-50%);z-index:0;pointer-events:none;animation:pulso 5s ease-in-out infinite}
.tierra-wrap{position:fixed;left:0;right:0;bottom:0;height:42vh;overflow:hidden;z-index:0;pointer-events:none}
.tierra{position:absolute;left:50%;top:63%;width:230vw;height:230vw;margin-left:-115vw;border-radius:50%;background:radial-gradient(circle at 50% 0%,#f6fbff 0%,#c9e8ff 3%,#6fb8ff 8%,#2f78e0 17%,#143f8c 36%,#081d4c 62%,#040c24 100%);box-shadow:0 -4px 22px 3px rgba(190,228,255,.95),0 -26px 90px 26px rgba(60,135,255,.5)}
.nube{position:absolute;border-radius:50%;background:radial-gradient(ellipse,rgba(255,255,255,.85),rgba(255,255,255,0) 70%);filter:blur(3px);opacity:.7}
.astro{position:fixed;z-index:1;pointer-events:none;filter:drop-shadow(0 0 22px rgba(90,160,255,.55))}
.astro svg{width:100%;height:auto;display:block}
.astro-gato{left:5vw;top:26vh;width:min(19vw,250px);animation:flotar 6.5s ease-in-out infinite}
.astro-perro{right:5vw;top:34vh;width:min(19vw,250px);animation:flotar2 7.5s ease-in-out infinite}
.login-titulo{position:relative;z-index:2;text-align:center;margin:8vh 0 .2rem}
.login-titulo h1{font-size:clamp(1.7rem,5.2vw,3.6rem);line-height:1.15;font-weight:800;margin:0;color:#e9f3ff;text-shadow:0 0 14px rgba(120,185,255,.85),0 0 34px rgba(60,120,255,.55);letter-spacing:.3px}
.login-titulo p{margin:.5rem 0 0;color:#b9d3ff;font-size:clamp(1.05rem,1.8vw,1.35rem);letter-spacing:.6px}
.login-sub{position:relative;z-index:2;text-align:center;color:#8fb2f0;margin:.2rem 0 1rem;font-size:.92rem}
.block-container{max-width:900px !important;margin-left:auto !important;margin-right:auto !important;position:relative;z-index:2;padding-top:2rem !important}
[data-testid="stForm"]{box-sizing:border-box;width:100%;max-width:440px;margin:0 auto !important;background:rgba(7,18,52,.58) !important;border:1px solid rgba(130,190,255,.5) !important;border-radius:20px !important;padding:22px 24px !important;box-shadow:0 0 44px rgba(60,130,255,.28),inset 0 0 34px rgba(90,160,255,.09);backdrop-filter:blur(7px)}
[data-testid="stForm"] label p,[data-testid="stForm"] label{color:#d5e6ff !important;font-weight:600}
[data-testid="stForm"] input{background:rgba(6,14,40,.85) !important;color:#9fe9ff !important;border:1px solid #3c63d8 !important;border-radius:12px !important}
[data-testid="stForm"] input:focus{border-color:#67e3ff !important;box-shadow:0 0 14px rgba(103,227,255,.6) !important}
[data-testid="stForm"] button{width:100%;background:linear-gradient(135deg,#6a5cf0 0%,#3a6bff 55%,#22c5e8 100%) !important;color:#fff !important;font-weight:700 !important;border:none !important;border-radius:12px !important;padding:.55rem 1rem !important;box-shadow:0 0 18px rgba(80,140,255,.55);transition:all .25s ease}
[data-testid="stForm"] button:hover{transform:translateY(-2px) scale(1.015);box-shadow:0 0 28px rgba(103,227,255,.8)}
[data-testid="stAlert"]{box-sizing:border-box;position:relative;z-index:2;width:100%;max-width:440px;margin:0 auto}
@media (max-width:900px){.astro-gato{left:2vw;top:1.5vh;width:72px}.astro-perro{right:2vw;top:1.5vh;width:72px}.login-titulo{margin-top:11vh}}
</style>"""
    css = css.replace("ESTRELLAS_A", estrellas_a).replace("ESTRELLAS_B", estrellas_b)
    estrella = """
<div class="estrella-mayor"><svg viewBox="-200 -200 400 400" xmlns="http://www.w3.org/2000/svg">
<defs>
<radialGradient id="brillo"><stop offset="0" stop-color="#ffffff"/><stop offset="0.05" stop-color="#f0f7ff"/><stop offset="0.16" stop-color="#9cc8ff" stop-opacity="0.75"/><stop offset="0.45" stop-color="#3a78ff" stop-opacity="0.28"/><stop offset="1" stop-color="#2050d0" stop-opacity="0"/></radialGradient>
<linearGradient id="rayo" gradientUnits="userSpaceOnUse" x1="0" y1="0" x2="0" y2="-190"><stop offset="0" stop-color="#ffffff" stop-opacity="0.95"/><stop offset="1" stop-color="#6aa8ff" stop-opacity="0"/></linearGradient>
<filter id="suave"><feGaussianBlur stdDeviation="1.6"/></filter>
</defs>
<circle r="200" fill="url(#brillo)"/>
<g filter="url(#suave)" stroke="url(#rayo)" stroke-linecap="round">
<line x1="0" y1="0" x2="0" y2="-190" stroke-width="5"/><line x1="0" y1="0" x2="0" y2="-190" stroke-width="5" transform="rotate(90)"/><line x1="0" y1="0" x2="0" y2="-190" stroke-width="5" transform="rotate(180)"/><line x1="0" y1="0" x2="0" y2="-190" stroke-width="5" transform="rotate(270)"/>
<line x1="0" y1="0" x2="0" y2="-120" stroke-width="3" transform="rotate(38)"/><line x1="0" y1="0" x2="0" y2="-120" stroke-width="3" transform="rotate(142)"/><line x1="0" y1="0" x2="0" y2="-120" stroke-width="3" transform="rotate(218)"/><line x1="0" y1="0" x2="0" y2="-120" stroke-width="3" transform="rotate(322)"/>
<line x1="0" y1="0" x2="0" y2="-150" stroke-width="2" transform="rotate(18)"/><line x1="0" y1="0" x2="0" y2="-150" stroke-width="2" transform="rotate(-18)"/>
</g>
</svg></div>"""
    tierra = """
<div class="tierra-wrap"><div class="tierra"></div>
<div class="nube" style="left:8%;bottom:5%;width:30%;height:5%"></div>
<div class="nube" style="left:40%;bottom:1%;width:34%;height:6%"></div>
<div class="nube" style="left:68%;bottom:7%;width:24%;height:4%"></div>
<div class="nube" style="left:22%;bottom:12%;width:22%;height:3%"></div></div>"""
    astros = f'<div class="astro astro-gato">{gato}</div><div class="astro astro-perro">{perro}</div>'
    titulo = """
<div class="login-titulo"><h1>🚀 Welcome Captain Larissa</h1><p>Gestión de Naves Espaciales</p></div>
<div class="login-sub">Ingresa la clave de acceso para despegar 🐾</div>"""
    cuerpo = estrella + tierra + astros + titulo
    # una sola línea: Streamlit interpreta las líneas en blanco o con sangría como texto, no como HTML
    return css + "".join(linea.strip() for linea in cuerpo.splitlines())


def obtener_password():
    clave = leer_secreto("APP_PASSWORD") or leer_secreto("password")
    return str(clave) if clave else None


def check_password():
    if st.session_state.get("password_correct"):
        return True
    clave = obtener_password()
    st.markdown(login_html(), unsafe_allow_html=True)
    if not clave:
        st.error("No hay contraseña configurada. Agrégala en Secrets como  APP_PASSWORD = \"tu-clave\".")
        return False
    with st.form("login"):
        ingreso = st.text_input("Contraseña", type="password")
        enviar = st.form_submit_button("🚀 Iniciar sesión")
    if enviar:
        if hmac.compare_digest(ingreso.encode("utf-8"), clave.encode("utf-8")):
            st.session_state["password_correct"] = True
            st.rerun()
        else:
            st.error("Contraseña incorrecta.")
    return False


if not check_password():
    st.stop()

@st.cache_resource
def _almacen_cacheado():
    return crear_almacen()


try:
    STORE = _almacen_cacheado()
    STORE.listar_fechas()  # prueba la conexión y que las tablas existan
except Exception as e:  # noqa: BLE001
    st.error("No se pudo conectar con la base de datos. Revisa SUPABASE_URL y SUPABASE_KEY en Secrets y que "
             "hayas ejecutado el archivo supabase_setup.sql en Supabase.")
    st.caption(f"Detalle técnico: {e}")
    st.stop()


@st.cache_data(ttl=120, show_spinner=False)
def _memo_cache(clave, version, _fn, _args):
    return _fn(*_args)


def _memo(clave, fn, *args):  # noqa: F811
    return _memo_cache((clave,) + tuple(str(a) for a in args), st.session_state.get("v_datos", 0), fn, args)


def _invalidar():  # noqa: F811
    st.session_state["v_datos"] = st.session_state.get("v_datos", 0) + 1


HOY = hoy_chile()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Limpieza automática de datos antiguos (se revisa una vez por sesión, al abrir la app)
if not st.session_state.get("limpieza_revisada"):
    st.session_state["limpieza_revisada"] = True
    try:
        _res_limpieza = ejecutar_limpieza(HOY)
        if _res_limpieza:
            st.session_state["aviso_limpieza"] = _res_limpieza
    except Exception as e:  # noqa: BLE001
        st.session_state["aviso_limpieza_error"] = str(e)


# =====================================================================
# AYUDAS DE INTERFAZ
# =====================================================================
def mostrar_df(obj):
    try:
        st.dataframe(obj, width="stretch")
    except Exception:  # noqa: BLE001
        st.dataframe(obj, use_container_width=True)


def estilizar(df, cols, rojo, verde):
    def f(v):
        if v == rojo:
            return "background-color: #ffcccc"
        if v == verde:
            return "background-color: #d4edda"
        return ""

    sty = df.style
    aplicar = getattr(sty, "map", None) or sty.applymap
    return aplicar(f, subset=cols)


def selector_rango(clave):
    c1, c2 = st.columns(2)
    f_ini = c1.date_input("Desde fecha", HOY, key=f"{clave}_ini")
    f_fin = c2.date_input("Hasta fecha", HOY, key=f"{clave}_fin")
    if f_ini > f_fin:
        st.warning("La fecha 'Desde' es posterior a 'Hasta'.")
        return None
    return f_ini, f_fin


@st.cache_data(show_spinner=False)
def _fecha_de_archivo(tipo, nombre, contenido):
    if tipo == "btk_yanez":
        try:
            return fecha_detectada(tipo, leer_archivo(nombre, contenido), nombre, HOY)
        except Exception:  # noqa: BLE001
            pass
    return fecha_desde_nombre(nombre, HOY)


def autodetectar_fecha(clave_uploader, clave_fecha, tipo):
    """Al subir un archivo, pone automáticamente la fecha que trae (la puedes cambiar después)."""
    f = st.session_state.get(clave_uploader)
    if f is None:
        return
    try:
        d = _fecha_de_archivo(tipo, f.name, f.getvalue())
    except Exception:  # noqa: BLE001
        d = None
    if d:
        st.session_state[clave_fecha] = d


def bloque_archivo(tipo, etiqueta, fecha, key, clave_fecha=None):
    detecta = clave_fecha is not None
    subido = st.file_uploader(f"Cargar {etiqueta} (Excel/CSV)", type=["xlsx", "xls", "csv"], key=key,
                              on_change=autodetectar_fecha if detecta else None,
                              args=(key, clave_fecha, tipo) if detecta else None)
    if detecta and subido is not None:
        try:
            d = _fecha_de_archivo(tipo, subido.name, subido.getvalue())
        except Exception:  # noqa: BLE001
            d = None
        if d and d != fecha:
            st.warning(f"📅 Este archivo parece ser del {d:%d/%m/%Y}, pero la fecha elegida es {fecha:%d/%m/%Y}. "
                       "Revisa que sea la correcta antes de procesar.")
        elif d:
            st.caption(f"📅 Fecha detectada: {d:%d/%m/%Y} ✔")
    guardado = nombre_archivo(tipo, fecha)
    if guardado and subido is None:
        st.caption(f"📂 Ya hay un archivo guardado para el {fecha:%d/%m/%Y}: {guardado} (se usa si no subes otro).")
    elif guardado and subido is not None:
        st.warning(f"Ya hay un archivo guardado para el {fecha:%d/%m/%Y} ({guardado}). "
                   "Al procesar se reemplazará por el que subiste.")
    return subido


def num(label, clave, fecha, previo):
    return st.number_input(label, min_value=0, value=int(previo.get(clave, 0)), step=1,
                           key=f"{clave}_{fecha.isoformat()}")


def aviso_no_entregados(nombre, tot, ent):
    if ent > tot:
        st.warning(f"{nombre}: los entregados ({ent}) superan a los totales ({tot}).")
    else:
        st.caption(f"No entregados {nombre}: {tot - ent}")


def alarma_hela_btk(df):
    """Alarma cuando Hela y BTK no coinciden en puntos o compromisos. El reporte se guía por el BTK."""
    flag = (df["U: Coincide puntos"].eq("NO") | df["W: Coincide compromisos"].eq("NO")
            | df["AB: Origen"].ne(ORIGEN_OK))
    dif = df[flag]
    if dif.empty:
        st.success("✅ Hela y BTK coinciden en todas las patentes.")
        return
    st.error(f"🚨 {len(dif)} de {len(df)} patentes no coinciden entre Hela y BTK. "
             "Los números del reporte salen del BTK; el Hela solo se usa como referencia.")
    cols = ["C: Patente", "B: Driver", "AB: Origen", "G: Total puntos HELA", "R: Puntos BTK",
            "AC: Dif. puntos (BTK − Hela)", "H: Compromisos HELA", "V: Compromisos BTK",
            "AD: Dif. compromisos (BTK − Hela)"]
    mostrar_df(dif[cols].reset_index(drop=True))
    st.caption("Si el BTK tiene más órdenes que el Hela, lo normal es que se hayan agregado durante el día, "
               "después de subir el Hela la noche anterior.")


def aviso_cambios_comuna(df):
    """Muestra las patentes donde la comuna pagada (según BTK) cambió respecto de Hela o se agregaron comunas."""
    cambio = df[df["Z: Coincide con Hela"].eq("NO")]
    agregadas = df[df["AE: Comunas agregadas desde BTK"].fillna("").ne("")]
    if not cambio.empty:
        st.warning(f"🔁 En {len(cambio)} patente(s) la comuna pagada cambió al usar el BTK.")
        mostrar_df(cambio[["C: Patente", "B: Driver", "Y: Comuna lejana según Hela", "AF: Valor según Hela",
                           "I: Comuna Lejana", "J: Valor"]].reset_index(drop=True))
    if not agregadas.empty:
        st.info(f"➕ Se agregaron comunas del BTK que no estaban en Hela en {len(agregadas)} patente(s).")
        mostrar_df(agregadas[["C: Patente", "B: Driver", "AE: Comunas agregadas desde BTK"]].reset_index(drop=True))


def aviso_rutas_mixtas(df, info):
    if info.get("modo_union") == "letra":
        st.info("🔗 El Hela no trae patentes: cada fila se tomó como una ruta (A, B, C… en el orden del archivo) y "
                "se unió con el BTK por la letra de la columna «conductor». Cada ruta quedó con una sola patente "
                "(la que llevó más pedidos con su letra) y los números salen de todos los pedidos de esa patente.")
    cruz = df[df["AH: Pedidos de otras rutas en esta patente"].fillna("").ne("")
              | df["AI: Pedidos de esta ruta llevados por otra patente"].fillna("").ne("")]
    if not cruz.empty:
        st.warning(f"🔀 {len(cruz)} ruta(s) tienen pedidos cruzados: una patente llevó pedidos con la letra de otra "
                   "ruta. Es lo que hace que los puntos del BTK cambien respecto del Hela.")
        mostrar_df(cruz[["AG: Letra de ruta", "C: Patente", "AJ: Pedidos con la letra de esta ruta (BTK)",
                         "AH: Pedidos de otras rutas en esta patente",
                         "AI: Pedidos de esta ruta llevados por otra patente"]].reset_index(drop=True))


@st.cache_data(show_spinner=False, max_entries=3)
def _leer_cache(nombre, contenido):
    return leer_archivo(nombre, contenido)


def autodetectar_mes():
    """Al subir el BTK del mes, pone automáticamente el año y el mes que trae la columna de fecha."""
    f = st.session_state.get("u_btk_mes")
    if f is None:
        return
    try:
        m = mes_detectado(_leer_cache(f.name, f.getvalue()))
    except Exception:  # noqa: BLE001
        m = None
    if m:
        st.session_state["m_anio"], st.session_state["m_mes"] = int(m[0]), int(m[1])


def resolver_mensual(tipo, primer, subido):
    """Guarda una copia del archivo solo si es liviano (los del mes pueden pesar varios MB)."""
    if subido is not None:
        contenido = subido.getvalue()
        if len(contenido) <= 1_500_000:
            guardar_archivo(tipo, primer, subido.name, contenido)
        else:
            st.caption(f"«{subido.name}» es grande: se procesa, pero no se guarda una copia del archivo "
                       "(sí se guarda el resultado).")
        return subido.name, contenido
    return archivo_guardado(tipo, primer)


def mostrar_mapeos(info):
    """Muestra qué columna del archivo se usó para cada dato y los avisos de contenido dudoso."""
    for clave, titulo in (("btk", "BTK Yáñez"), ("hela", "Reporte Hela"), ("paris", "BTK París")):
        for aviso in info.get(f"avisos_{clave}", []):
            st.warning(f"{titulo}: {aviso}")
        filas = info.get(f"mapeo_{clave}")
        if filas:
            with st.expander(f"🧭 Columnas usadas del {titulo}"):
                mostrar_df(pd.DataFrame(filas))


def mostrar_ns(m, dfs):
    c = calcular_ns(m)
    fecha = datetime.date.fromisoformat(m["fecha"])
    st.markdown("---")
    st.subheader(f"📈 Resultados NS del {fecha:%d/%m/%Y}")
    mostrar_mapeos(m)

    def fila(titulo, bloque):
        cols = st.columns(3)
        for col, (nombre, k) in zip(cols, (("General", "gen"), ("Easy", "easy"), ("París", "paris"))):
            col.metric(f"{titulo} {nombre}", fmt_pct(c[f"ns_{bloque}_{k}"]))
            col.caption(f"Total: {c[f'{bloque}_{k}_tot']} | Entregados: {c[f'{bloque}_{k}_ent']} | "
                        f"No entregados: {c[f'{bloque}_{k}_no']}")

    fila("NS Compromiso", "fc")
    fila("NS Ácido", "acid")
    fila("NS Retiros", "ret")

    c1, c2, c3 = st.columns(3)
    c1.metric("Rutas Total Móviles", c["moviles_gen"])
    c2.metric("Móviles Easy", c["moviles_easy"])
    c3.metric("Móviles París", c["moviles_paris"])

    if "fc_btk_easy_tot" in m:
        ref = pd.DataFrame([
            {"Cliente": "Easy", "FC manual": c["fc_easy_tot"], "FC según BTK": m["fc_btk_easy_tot"],
             "Entregados manual": c["fc_easy_ent"], "Entregados según BTK": m["fc_btk_easy_ent"]},
            {"Cliente": "París", "FC manual": c["fc_paris_tot"], "FC según BTK": m["fc_btk_paris_tot"],
             "Entregados manual": c["fc_paris_ent"], "Entregados según BTK": m["fc_btk_paris_ent"]},
        ])
        ref["Revisión"] = ["OK" if r["FC manual"] == r["FC según BTK"] and
                           r["Entregados manual"] == r["Entregados según BTK"] else "⚠️ Revisar"
                           for _, r in ref.iterrows()]
        with st.expander("🔎 Verificación automática con el BTK (órdenes cuya fecha compromiso = fecha del informe)"):
            mostrar_df(ref)
            st.caption("Es solo una referencia para detectar errores de digitación; el reporte usa tus datos manuales.")

    st.subheader("⚠️ Sub-estados de No Entrega del Día")
    sub = dfs.get("subestados", pd.DataFrame())
    if sub.empty:
        st.info("No hubo pedidos no entregados.")
    else:
        mostrar_df(sub)

    st.subheader("📥 Descargas")
    d1, d2 = st.columns(2)
    hojas = {
        "Resumen": resumen_ns_df(m),
        "Sub-estados No Entrega": sub,
        "Detalle No Entregados": dfs.get("detalle_no_entregados", pd.DataFrame()),
        "Detalle Easy y París": dfs.get("detalle", pd.DataFrame()),
    }
    d1.download_button("📥 Descargar Excel (con detalle)", a_excel(hojas),
                       file_name=f"NS_Diario_{fecha.isoformat()}.xlsx", mime=XLSX, key=f"dl_xlsx_ns_{fecha}")
    d2.download_button("📄 Descargar informe PDF", pdf_ns(m, sub),
                       file_name=f"informe_nivel_servicio_{fecha.isoformat()}.pdf", mime="application/pdf",
                       key=f"dl_pdf_ns_{fecha}")


def mostrar_semanal(f_ini, f_fin):
    datos = cargar("ns", f_ini, f_fin)
    if not datos:
        st.info("No hay reportes NS diarios guardados en ese rango. Procésalos primero en la pestaña 1.")
        return

    con_datos = {f for f, _, _ in datos}
    faltan, d = [], f_ini
    while d <= f_fin:
        if d.weekday() < 6 and d not in con_datos:
            faltan.append(d.strftime("%d/%m"))
        d += datetime.timedelta(days=1)
    if faltan:
        st.warning("Días de lunes a sábado sin reporte guardado: " + ", ".join(faltan))
    if any(f.weekday() == 6 for f in con_datos):
        st.info("El rango incluye días domingo con datos.")

    st.subheader("📅 NS Fecha Compromiso por día")
    mostrar_df(tabla_fc(datos, texto=True))

    r = resumen_periodo(datos)
    st.subheader("🌐 Métricas Globales del Rango Seleccionado")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Órdenes Salidas a Ruta", r["ordenes"])
    m2.metric("Órdenes Entregadas", r["entregadas"])
    m3.metric("Órdenes No Entregadas", r["no_entregadas"])
    m4.metric("NS Ácido General", fmt_pct(r["ns_acid_gen"]))
    n1, n2, n3 = st.columns(3)
    n1.metric("NS Ácido Easy", fmt_pct(r["ns_acid_easy"]))
    n2.metric("NS Ácido París", fmt_pct(r["ns_acid_paris"]))
    n3.metric("% No Entregados (ácido)", fmt_pct(r["pct_no_ent"]))

    st.markdown("---")
    st.subheader("🚛 Promedio y Totales de Flota")
    f1, f2, f3 = st.columns(3)
    f1.metric("Promedio Móviles Diarios en Ruta", f"{r['prom_moviles']:.1f}")
    f2.metric("Rutas Operadas Easy", r["rutas_easy"])
    f3.metric("Rutas Operadas París", r["rutas_paris"])

    dias = df_ns_dias(datos)
    st.subheader("📈 Evolución diaria")
    st.line_chart(dias.set_index("Fecha")[["NS FC General (%)", "NS Ácido General (%)"]])

    st.markdown("---")
    st.subheader("⚠️ Consolidado de Submotivos de No Entrega")
    rank = ranking_submotivos(datos)
    if rank.empty:
        st.info("No se registraron no entregas en los días seleccionados.")
    else:
        mostrar_df(rank)

    resumen = pd.DataFrame([
        {"Indicador": "Órdenes salidas a ruta", "Valor": r["ordenes"]},
        {"Indicador": "Órdenes entregadas", "Valor": r["entregadas"]},
        {"Indicador": "Órdenes no entregadas", "Valor": r["no_entregadas"]},
        {"Indicador": "NS Ácido general (%)", "Valor": round(r["ns_acid_gen"], 2)},
        {"Indicador": "NS Ácido Easy (%)", "Valor": round(r["ns_acid_easy"], 2)},
        {"Indicador": "NS Ácido París (%)", "Valor": round(r["ns_acid_paris"], 2)},
        {"Indicador": "% No entregados (ácido)", "Valor": round(r["pct_no_ent"], 2)},
        {"Indicador": "Promedio móviles diarios", "Valor": round(r["prom_moviles"], 1)},
        {"Indicador": "Rutas operadas Easy", "Valor": r["rutas_easy"]},
        {"Indicador": "Rutas operadas París", "Valor": r["rutas_paris"]},
    ])
    xls = a_excel({"NS FC por fecha": tabla_fc(datos), "Resumen": resumen,
                   "Submotivos": rank, "Detalle por día": dias})
    st.download_button("📥 Descargar reporte semanal en Excel", xls,
                       file_name=f"Reporte_Semanal_{f_ini.isoformat()}_a_{f_fin.isoformat()}.xlsx",
                       mime=XLSX, key="dl_sem")


# =====================================================================
# PÁGINA
# =====================================================================
st.title("🚚 Suite de Reportes Operativos - Transportes Yáñez")
if "aviso_limpieza" in st.session_state:
    _corte, _na, _nr = st.session_state["aviso_limpieza"]
    st.info(f"🧹 Limpieza automática: se borró todo lo anterior al {_corte:%d/%m/%Y} "
            f"({_na} archivo(s) y {_nr} reporte(s)).")
if "aviso_limpieza_error" in st.session_state:
    st.warning("No se pudo revisar la limpieza automática: " + st.session_state["aviso_limpieza_error"])
_prox_limpieza = proxima_limpieza(leer_config_limpieza(), HOY)
if _prox_limpieza and 0 <= (_prox_limpieza[0] - HOY).days <= 7:
    st.warning(f"📅 El {_prox_limpieza[0]:%d/%m/%Y} se borrarán todos los datos anteriores al "
               f"{_prox_limpieza[1]:%d/%m/%Y}. Antes de esa fecha descarga tus reportes del mes y una copia de "
               "seguridad (pestaña 6).")
if st.sidebar.button("Cerrar sesión"):
    st.session_state["password_correct"] = False
    st.rerun()
if STORE.nombre == "supabase":
    st.sidebar.success("Datos guardados en Supabase ✔")
else:
    st.sidebar.warning("Guardando en un archivo local: se pierde si la app se reinicia. "
                       "Configura SUPABASE_URL y SUPABASE_KEY en Secrets.")
st.sidebar.caption("Cada archivo que subes queda registrado con la fecha que elegiste.")
st.sidebar.caption(f"Versión de la app: {APP_VERSION}")

tabs = st.tabs([
    "1. NS Diario (Easy & París)",
    "2. Reporte Comunas",
    "3. Reporte Bonificación",
    "4. Cruce de Calces París",
    "5. Consolidado Semanal/Rango",
    "6. Descargas y Respaldo",
    "7. Reporte Mensual",
])

# ---------------------------------------------------------------------
# TAB 1: NS DIARIO
# ---------------------------------------------------------------------
with tabs[0]:
    st.header("1. Reporte Nivel de Servicio (NS) Diario")
    fecha_ns = st.date_input("Fecha a la que corresponde este archivo BTK", HOY, key="f_ns")
    previo_ns = cargar("ns", fecha_ns, fecha_ns)
    prev_meta = previo_ns[0][1] if previo_ns else {}

    up_btk = bloque_archivo("btk_yanez", "Archivo BTK YÁÑEZ", fecha_ns, "u_btk_main", "f_ns")

    st.subheader("📊 Datos Fecha Compromiso y Retiros (ingreso manual)")
    col_e, col_p = st.columns(2)
    with col_e:
        st.markdown("**EASY - Fecha Compromiso**")
        fc_easy_tot = num("Totales FC Easy", "fc_easy_tot", fecha_ns, prev_meta)
        fc_easy_ent = num("Entregados FC Easy", "fc_easy_ent", fecha_ns, prev_meta)
        aviso_no_entregados("FC Easy", fc_easy_tot, fc_easy_ent)
        st.markdown("**EASY - Retiros**")
        ret_easy_tot = num("Totales Retiros Easy", "ret_easy_tot", fecha_ns, prev_meta)
        ret_easy_ent = num("Entregados / recogidos Retiros Easy", "ret_easy_ent", fecha_ns, prev_meta)
        aviso_no_entregados("Retiros Easy", ret_easy_tot, ret_easy_ent)
    with col_p:
        st.markdown("**PARÍS - Fecha Compromiso**")
        fc_paris_tot = num("Totales FC París", "fc_paris_tot", fecha_ns, prev_meta)
        fc_paris_ent = num("Entregados FC París", "fc_paris_ent", fecha_ns, prev_meta)
        aviso_no_entregados("FC París", fc_paris_tot, fc_paris_ent)
        st.markdown("**PARÍS - Retiros / Recogidos**")
        ret_paris_tot = num("Totales Retiros París", "ret_paris_tot", fecha_ns, prev_meta)
        ret_paris_ent = num("Recogidos Retiros París", "ret_paris_ent", fecha_ns, prev_meta)
        aviso_no_entregados("Retiros París (no recogidos)", ret_paris_tot, ret_paris_ent)

    if st.button("Procesar y Guardar NS Diario", key="b_proc_ns"):
        arch = resolver_archivo("btk_yanez", fecha_ns, up_btk)
        if arch is None:
            st.error("Sube el archivo BTK YÁÑEZ (o guarda uno antes para esta fecha).")
        else:
            try:
                b = preparar_btk(leer_archivo(*arch))
                manual = {
                    "fc_easy_tot": fc_easy_tot, "fc_easy_ent": fc_easy_ent,
                    "fc_paris_tot": fc_paris_tot, "fc_paris_ent": fc_paris_ent,
                    "ret_easy_tot": ret_easy_tot, "ret_easy_ent": ret_easy_ent,
                    "ret_paris_tot": ret_paris_tot, "ret_paris_ent": ret_paris_ent,
                }
                meta, dfs = procesar_ns(b, fecha_ns, manual)
                meta["archivo_btk"] = arch[0]
                guardar_resultado("ns", fecha_ns, meta, dfs)
                st.success(f"¡Día {fecha_ns:%d/%m/%Y} procesado y guardado!")
            except Exception as e:  # noqa: BLE001
                st.error(f"No se pudo procesar el archivo: {e}")

    res_ns = cargar("ns", fecha_ns, fecha_ns)
    if res_ns:
        mostrar_ns(res_ns[0][1], res_ns[0][2])
    else:
        st.info("Todavía no hay un reporte guardado para esta fecha.")

# ---------------------------------------------------------------------
# TAB 2: COMUNAS
# ---------------------------------------------------------------------
with tabs[1]:
    st.header("2. Reporte de Comunas")
    fecha_com = st.date_input("Fecha a la que corresponden estos archivos", HOY, key="f_com_arch")
    cc1, cc2 = st.columns(2)
    with cc1:
        up_hela = bloque_archivo("hela", "Reporte HELA", fecha_com, "u_hela", "f_com_arch")
    with cc2:
        up_btk_com = bloque_archivo("btk_yanez", "BTK YÁÑEZ", fecha_com, "u_btk_com", "f_com_arch")

    if st.button("Procesar y Guardar Comunas del Día", key="b_proc_com"):
        a_hela = resolver_archivo("hela", fecha_com, up_hela)
        a_btk = resolver_archivo("btk_yanez", fecha_com, up_btk_com)
        if a_hela is None or a_btk is None:
            st.error("Faltan archivos: necesitas el Reporte HELA y el BTK YÁÑEZ de esa fecha.")
        else:
            try:
                b = preparar_btk(leer_archivo(*a_btk))
                df_res, info = procesar_comunas(leer_archivo(*a_hela), b, fecha_com)
                guardar_resultado("comunas", fecha_com, info, {"reporte": df_res})
                st.success(f"¡Reporte de Comunas del {fecha_com:%d/%m/%Y} procesado y guardado!")
                mostrar_mapeos(info)
                aviso_rutas_mixtas(df_res, info)
                alarma_hela_btk(df_res)
                aviso_cambios_comuna(df_res)
            except Exception as e:  # noqa: BLE001
                st.error(f"No se pudo procesar: {e}")

    st.markdown("---")
    st.subheader("🔍 Consultar Comunas por Rango de Fechas")
    rango_c = selector_rango("com")
    if rango_c:
        datos_c = cargar("comunas", *rango_c)
        df_com = concat_dfs(datos_c, "reporte")
        if df_com.empty:
            st.warning("No hay registros de comunas en el rango seleccionado.")
        else:
            sin_matriz = df_com["AA: Comunas no reconocidas"].fillna("").ne("").sum()
            malos = (df_com["U: Coincide puntos"].eq("NO") | df_com["W: Coincide compromisos"].eq("NO")
                     | df_com["AB: Origen"].ne(ORIGEN_OK)).sum()
            k1, k2, k3 = st.columns(3)
            k1.metric("Filas (patente-día)", len(df_com))
            k2.metric("Patentes con diferencias Hela vs BTK", int(malos))
            k3.metric("Con comunas no reconocidas", int(sin_matriz))
            mostrar_df(estilizar(df_com, ["U: Coincide puntos", "W: Coincide compromisos", "Z: Coincide con Hela"],
                                 "NO", "SÍ"))
            st.download_button("📥 Descargar Reporte de Comunas en Excel", a_excel({"Comunas": df_com}),
                               file_name=f"Comunas_{rango_c[0].isoformat()}_a_{rango_c[1].isoformat()}.xlsx",
                               mime=XLSX, key="dl_com")

# ---------------------------------------------------------------------
# TAB 3: BONIFICACIÓN
# ---------------------------------------------------------------------
with tabs[2]:
    st.header("3. Reporte de Bonificación")
    st.markdown("Se arma automáticamente desde los Reportes de Comunas guardados en el rango que elijas.")
    rango_b = selector_rango("bon")
    excluir = st.checkbox("Excluir patentes sin ruta en el BTK (no salieron a ruta)", value=True, key="bon_excl")
    if rango_b:
        df_c_src = concat_dfs(cargar("comunas", *rango_b), "reporte")
        if df_c_src.empty:
            st.warning("No hay reportes de comunas guardados en ese rango. Procésalos primero en la pestaña 2.")
        else:
            df_bonif = df_bonificacion(df_c_src, excluir)
            por_patente = bonificacion_por_patente(df_bonif)
            omitidas = len(df_bonificacion(df_c_src, excluir, incluir_cero=True)) - len(df_bonif)
            b1, b2, b3 = st.columns(3)
            b1.metric("Total bonificación del rango", f"$ {int(df_bonif['D: Valor'].sum()):,}".replace(",", "."))
            b2.metric("Rutas con pago", len(df_bonif))
            b3.metric("Patentes con bonificación", len(por_patente))
            st.caption(f"Solo se listan las rutas con pago. Se omitieron {omitidas} ruta(s) con $0 "
                       "(no llevaron ninguna comuna con valor).")
            if df_bonif.empty:
                st.info("Ninguna ruta tuvo bonificación en este rango.")
            mostrar_df(df_bonif)
            st.subheader("Total por patente")
            mostrar_df(por_patente)
            st.download_button("📥 Descargar Bonificación en Excel",
                               a_excel({"Bonificación": df_bonif, "Total por patente": por_patente}),
                               file_name=f"Bonificacion_{rango_b[0].isoformat()}_a_{rango_b[1].isoformat()}.xlsx",
                               mime=XLSX, key="dl_bon")

# ---------------------------------------------------------------------
# TAB 4: CALCES PARÍS
# ---------------------------------------------------------------------
with tabs[3]:
    st.header("4. Reporte Cruce de Calces París")
    fecha_cal = st.date_input("Fecha a la que corresponden estos archivos", HOY, key="f_calce_arch")
    cy, cp = st.columns(2)
    with cy:
        up_y = bloque_archivo("btk_yanez", "BTK YÁÑEZ", fecha_cal, "u_y_c", "f_calce_arch")
    with cp:
        up_p = bloque_archivo("btk_paris", "BTK PARÍS", fecha_cal, "u_p_c")

    if st.button("Procesar y Guardar Calces del Día", key="b_proc_calce"):
        a_y = resolver_archivo("btk_yanez", fecha_cal, up_y)
        a_p = resolver_archivo("btk_paris", fecha_cal, up_p)
        if a_y is None or a_p is None:
            st.error("Faltan archivos: necesitas el BTK YÁÑEZ y el BTK PARÍS de esa fecha.")
        else:
            try:
                y = preparar_btk(leer_archivo(*a_y))
                df_cal, info = procesar_calces(y, leer_archivo(*a_p), fecha_cal)
                guardar_resultado("calces", fecha_cal, info, {"cruce": df_cal})
                st.success(f"¡Cruce de Calces del {fecha_cal:%d/%m/%Y} guardado!")
                mostrar_mapeos(info)
                if info["guias_duplicadas_paris"]:
                    st.warning(f"El BTK París tiene {info['guias_duplicadas_paris']} guías repetidas; "
                               "se usó la primera aparición de cada una.")
            except Exception as e:  # noqa: BLE001
                st.error(f"No se pudo procesar: {e}")

    st.markdown("---")
    st.subheader("🔍 Consultar Cruce de Calces por Rango de Fechas")
    rango_cal = selector_rango("cal")
    if rango_cal:
        df_cal_all = concat_dfs(cargar("calces", *rango_cal), "cruce")
        if df_cal_all.empty:
            st.warning("No hay registros de calces en el rango seleccionado.")
        else:
            total = len(df_cal_all)
            no_coin = int((df_cal_all["Estado Cruce"] == "No Coincide").sum())
            q1, q2, q3 = st.columns(3)
            q1.metric("Guías cruzadas", total)
            q2.metric("No coinciden", no_coin)
            q3.metric("% Coincidencia", fmt_pct(pct(total - no_coin, total)))

            solo_no = st.checkbox("Ver solo las que NO coinciden", key="cal_solo_no")
            vista = df_cal_all[df_cal_all["Estado Cruce"] == "No Coincide"] if solo_no else df_cal_all
            mostrar_df(estilizar(vista.reset_index(drop=True), ["Estado Cruce"], "No Coincide", "Coincide"))

            if no_coin:
                with st.expander("Combinaciones de sub-estados que no coinciden (para revisar las reglas)"):
                    comb = (df_cal_all[df_cal_all["Estado Cruce"] == "No Coincide"]
                            .groupby(["Beetrack T.Yañez Estado", "Beetrack T.Yañez Sub Estado", "Beetrack París", "Beetrack París Sub Estado", "Diagnóstico automático"])
                            .size().reset_index(name="Cantidad").sort_values("Cantidad", ascending=False))
                    mostrar_df(comb)
            st.download_button("📥 Descargar Cruce de Calces en Excel",
                               a_excel({"Calces París": df_cal_all,
                                        "No coinciden": df_cal_all[df_cal_all["Estado Cruce"] == "No Coincide"]}),
                               file_name=f"Calces_Paris_{rango_cal[0].isoformat()}_a_{rango_cal[1].isoformat()}.xlsx",
                               mime=XLSX, key="dl_cal")

# ---------------------------------------------------------------------
# TAB 5: SEMANAL / RANGO
# ---------------------------------------------------------------------
with tabs[4]:
    st.header("5. Reporte Semanal y Rango de Fechas (Lunes a Sábado)")
    modo = st.radio("Tipo de consulta", ["Semana (lunes a sábado)", "Rango libre"], horizontal=True, key="modo_sem")
    rango_s = None
    if modo.startswith("Semana"):
        dia = st.date_input("Elige cualquier día de la semana", HOY, key="dia_sem")
        lunes = dia - datetime.timedelta(days=dia.weekday())
        sabado = lunes + datetime.timedelta(days=5)
        st.caption(f"Semana del {lunes:%d/%m/%Y} (lunes) al {sabado:%d/%m/%Y} (sábado)")
        rango_s = (lunes, sabado)
    else:
        rango_s = selector_rango("sem")
    if rango_s:
        mostrar_semanal(*rango_s)

# ---------------------------------------------------------------------
# TAB 6: DESCARGAS Y RESPALDO
# ---------------------------------------------------------------------
with tabs[5]:
    st.header("6. Descargas y Respaldo")

    st.subheader("📦 Excel con todos los reportes")
    st.markdown("Junta NS diario, sub-estados, comunas, bonificación y calces del rango en un solo Excel.")
    rango_t = selector_rango("todo")
    if rango_t and st.button("Preparar Excel completo", key="b_todo"):
        xls = excel_completo(*rango_t)
        if xls is None:
            st.warning("No hay datos guardados en ese rango.")
            st.session_state.pop("excel_todo", None)
        else:
            st.session_state["excel_todo"] = (xls, f"Reportes_{rango_t[0].isoformat()}_a_{rango_t[1].isoformat()}.xlsx")
    if "excel_todo" in st.session_state:
        st.download_button("📥 Descargar Excel completo", st.session_state["excel_todo"][0],
                           file_name=st.session_state["excel_todo"][1], mime=XLSX, key="dl_todo")

    st.markdown("---")
    st.subheader("🗂️ Archivos guardados")
    arch_df = listar_archivos()
    if arch_df.empty:
        st.info("Todavía no hay archivos guardados.")
    else:
        mostrar_df(arch_df)
    fechas_guardadas = listar_fechas()
    if fechas_guardadas:
        col_b1, col_b2 = st.columns([2, 1])
        f_borrar = col_b1.selectbox("Fecha a eliminar (archivos y reportes de ese día)", fechas_guardadas, key="f_borrar")
        confirma = col_b2.checkbox("Confirmo que quiero borrarla", key="c_borrar")
        if st.button("Eliminar fecha", key="b_borrar", disabled=not confirma):
            borrar_fecha(f_borrar)
            st.success(f"Se eliminó todo lo guardado del {f_borrar}.")
            st.rerun()

    st.markdown("---")
    st.subheader("🧹 Limpieza automática de datos antiguos")
    cfg_l = leer_config_limpieza()
    cl1, cl2, cl3 = st.columns(3)
    activa_l = cl1.checkbox("Activar limpieza automática", value=bool(cfg_l["activa"]), key="lim_activa")
    dia_l = int(cl2.number_input("Día del mes en que se limpia", min_value=1, max_value=28, step=1,
                                 value=int(cfg_l["dia"]), key="lim_dia"))
    prim_l = cl3.date_input("No limpiar antes del", datetime.date.fromisoformat(cfg_l["primera"]), key="lim_primera")
    if st.button("Guardar configuración de limpieza", key="b_lim_cfg"):
        guardar_config_limpieza({**cfg_l, "activa": activa_l, "dia": dia_l, "primera": prim_l.isoformat()})
        st.success("Configuración guardada.")
        st.rerun()
    prox_l = proxima_limpieza(cfg_l, HOY)
    if prox_l:
        st.info(f"Próxima limpieza: **{prox_l[0]:%d/%m/%Y}**. Se borrará todo lo anterior al {prox_l[1]:%d/%m/%Y} "
                f"(hasta el {prox_l[1] - datetime.timedelta(days=1):%d/%m/%Y}). Se ejecuta la primera vez que "
                "abras la app ese día o después; la app no corre sola.")
    else:
        st.caption("La limpieza automática está desactivada.")
    if cfg_l.get("ultima"):
        st.caption(f"Última limpieza: {datetime.date.fromisoformat(cfg_l['ultima']):%d/%m/%Y}")
    with st.expander("Limpiar ahora"):
        corte_m = st.date_input("Borrar todo lo anterior al", value=datetime.date(HOY.year, HOY.month, 1),
                                key="lim_corte")
        ok_m = st.checkbox("Ya descargué mis reportes y una copia de seguridad", key="lim_ok")
        if st.button("Borrar ahora", key="b_lim_ya", disabled=not ok_m):
            na, nr = limpiar_antes_de(corte_m)
            st.success(f"Se borraron {na} archivo(s) y {nr} reporte(s) anteriores al {corte_m:%d/%m/%Y}.")
            st.rerun()

    st.markdown("---")
    st.subheader("💾 Copia de seguridad")
    st.markdown(
        "Descarga una copia de todo lo guardado (archivos y reportes) en un solo archivo `.json.gz`. "
        "Sirve como seguridad extra o para mover los datos a otra base."
    )
    if st.button("Preparar copia de seguridad", key="b_resp_prep"):
        with st.spinner("Preparando copia..."):
            st.session_state["respaldo"] = (exportar_respaldo(), f"respaldo_yanez_{HOY.isoformat()}.json.gz")
    if "respaldo" in st.session_state:
        st.download_button("📥 Descargar copia de seguridad", st.session_state["respaldo"][0],
                           file_name=st.session_state["respaldo"][1], mime="application/gzip", key="dl_respaldo")
    up_resp = st.file_uploader("Restaurar desde una copia (.json.gz)", type=["gz"], key="u_resp")
    if up_resp is not None and st.button("Restaurar copia", key="b_resp"):
        try:
            n_a, n_r = importar_respaldo(up_resp.getvalue())
            st.success(f"Copia restaurada: {n_a} archivos y {n_r} reportes.")
        except Exception as e:  # noqa: BLE001
            st.error(f"No se pudo restaurar: {e}")


# ---------------------------------------------------------------------
# TAB 7: REPORTE MENSUAL
# ---------------------------------------------------------------------
with tabs[6]:
    st.header("7. Reporte Mensual")
    st.markdown("Sube el BTK y el Hela de todo el mes (con la **fecha de cada día en la columna A**) para obtener "
                "comunas, NS compromiso, NS ácido y órdenes por patente del mes.")
    st.session_state.setdefault("m_anio", HOY.year)
    st.session_state.setdefault("m_mes", HOY.month)
    cm1, cm2 = st.columns(2)
    anio_m = int(cm1.number_input("Año", min_value=2020, max_value=2100, step=1, key="m_anio"))
    mes_m = int(cm2.selectbox("Mes", list(range(1, 13)), format_func=lambda i: MESES[i - 1].capitalize(), key="m_mes"))
    primer_m = primer_dia_mes(anio_m, mes_m)
    ultimo_m = (primer_m + datetime.timedelta(days=32)).replace(day=1) - datetime.timedelta(days=1)

    st.subheader("📂 Archivos del mes")
    cu1, cu2, cu3 = st.columns(3)
    up_btk_m = cu1.file_uploader("BTK Yáñez del mes (Excel/CSV)", type=["xlsx", "xls", "csv"], key="u_btk_mes",
                                 on_change=autodetectar_mes)
    up_hela_m = cu2.file_uploader("Hela del mes (Excel/CSV)", type=["xlsx", "xls", "csv"], key="u_hela_mes")
    up_paris_m = cu3.file_uploader("BTK París del mes (solo para el calce)", type=["xlsx", "xls", "csv"],
                                   key="u_paris_mes")
    g_btk, g_hela = nombre_archivo("btk_mensual", primer_m), nombre_archivo("hela_mensual", primer_m)
    g_par = nombre_archivo("btk_paris_mensual", primer_m)
    if (g_btk or g_hela or g_par) and up_btk_m is None and up_hela_m is None and up_paris_m is None:
        st.caption(f"📂 Guardado para {MESES[mes_m - 1]} de {anio_m}: BTK Yáñez «{g_btk or '—'}», "
                   f"Hela «{g_hela or '—'}», BTK París «{g_par or '—'}».")

    if st.button("Procesar mes", key="b_proc_mes"):
        a_btk = resolver_mensual("btk_mensual", primer_m, up_btk_m)
        a_hela = resolver_mensual("hela_mensual", primer_m, up_hela_m)
        if a_btk is None or a_hela is None:
            st.error("Faltan archivos: necesitas el BTK y el Hela del mes.")
        else:
            try:
                with st.spinner("Procesando el mes completo…"):
                    b_mes = preparar_btk(_leer_cache(*a_btk))
                    meta_p, dfs_p = procesar_mensual(_leer_cache(*a_hela), b_mes, anio_m, mes_m)
                    guardar_resultado("mensual", primer_m, meta_p, dfs_p)
                st.success(f"¡{MESES[mes_m - 1].capitalize()} de {anio_m} procesado! {meta_p['dias']} día(s) con datos.")
                for aviso in meta_p["avisos"]:
                    st.warning(aviso)
                mostrar_mapeos(meta_p)
            except Exception as e:  # noqa: BLE001
                st.error(f"No se pudo procesar el mes: {e}")

    res_mes = cargar("mensual", primer_m, primer_m)
    meta_m, dfs_m = (res_mes[0][1], res_mes[0][2]) if res_mes else ({}, {})
    regs_m = cargar_fc_mes(primer_m, ultimo_m)
    res_cal = cargar("calce_mensual", primer_m, primer_m)
    meta_cal, dfs_cal = (res_cal[0][1], res_cal[0][2]) if res_cal else ({}, {})
    sub = st.tabs(["Comunas del mes", "NS Compromiso", "NS Ácido", "NS por patente (mes)", "Descargar todo",
                   "Calce París del mes"])
    sin_datos = "Todavía no hay datos de este mes. Sube los archivos y presiona «Procesar mes»."

    # ----- Comunas del mes -----
    with sub[0]:
        com_m = dfs_m.get("comunas", pd.DataFrame())
        if com_m.empty:
            st.info(sin_datos)
        else:
            for aviso in meta_m.get("avisos", []):
                st.warning(aviso)
            bon_m = df_bonificacion(com_m)
            q1, q2, q3, q4 = st.columns(4)
            q1.metric("Días con datos", meta_m.get("dias", 0))
            q2.metric("Rutas-día", len(com_m))
            q3.metric("Rutas con pago", len(bon_m))
            q4.metric("Bonificación del mes", f"$ {int(bon_m['D: Valor'].sum()):,}".replace(",", "."))
            mostrar_df(estilizar(com_m, ["U: Coincide puntos", "W: Coincide compromisos", "Z: Coincide con Hela"],
                                 "NO", "SÍ"))
            st.subheader("Bonificación del mes por patente")
            mostrar_df(bonificacion_por_patente(bon_m))
            pat_c = dfs_m.get("patentes", pd.DataFrame())
            if not pat_c.empty:
                st.subheader("NS de cada patente en el mes")
                st.caption("Resumen del mes por patente, aparte del detalle de cada día de la tabla de arriba.")
                mostrar_df(pat_c)
            hojas_c = {"Comunas del mes": com_m, "Bonificación del mes": bon_m,
                       "Bonificación por patente": bonificacion_por_patente(bon_m)}
            if not pat_c.empty:
                hojas_c["NS por patente del mes"] = pat_c
            dm1, dm2 = st.columns(2)
            dm1.download_button("📥 Descargar comunas y bonificación del mes (Excel)",
                                a_excel(hojas_c),
                                file_name=f"Comunas_{anio_m}-{mes_m:02d}.xlsx", mime=XLSX, key="dl_com_mes")
            dm2.download_button("📄 Descargar bonificación del mes en PDF", pdf_bonificacion(com_m, anio_m, mes_m),
                                file_name=f"Bonificacion_{anio_m}-{mes_m:02d}.pdf", mime="application/pdf",
                                key="dl_bon_pdf")

    # ----- NS Compromiso (ingreso manual) -----
    with sub[1]:
        st.subheader("Datos de fecha compromiso por día (ingreso manual)")
        st.caption("Fecha compromiso = pedidos cuya fecha de entrega límite es ese mismo día, aunque se hayan "
                   "entregado después (incluso pedidos del mes anterior). Por eso se ingresan a mano, por cliente.")
        dia_fc = st.date_input("Día que estás ingresando", value=min(max(HOY, primer_m), ultimo_m),
                               min_value=primer_m, max_value=ultimo_m, key=f"m_fc_dia_{anio_m}_{mes_m}")
        ya = {r["fecha"]: r for r in regs_m}
        pre_e = pre_p = (0, 0, 0)
        origen_pre = None
        if dia_fc in ya and ya[dia_fc]["easy"] is not None:
            pre_e, pre_p = ya[dia_fc]["easy"], ya[dia_fc]["paris"]
            if ya[dia_fc].get("origen") == "NS diario":
                origen_pre = "ns_diario"
                st.caption("Este día toma los datos del NS diario (pestaña 1). Si los guardas aquí, se usará lo de aquí.")
            else:
                origen_pre = "guardado"
        elif dia_fc in ya:
            origen_pre = "general"
            st.warning(f"Este día solo tiene un total general guardado ({ya[dia_fc]['tot']} pedidos). "
                       "Ingrésalo ahora por cliente para reemplazarlo.")
        else:
            ns_dia = cargar("ns", dia_fc, dia_fc)
            if ns_dia:
                cc = calcular_ns(ns_dia[0][1])
                if cc["fc_gen_tot"]:
                    pre_e = (cc["fc_easy_tot"], cc["fc_easy_ent"], cc["fc_easy_no"])
                    pre_p = (cc["fc_paris_tot"], cc["fc_paris_ent"], cc["fc_paris_no"])
                    origen_pre = "ns_diario"
                    st.caption("Valores precargados desde el NS diario de ese día; puedes cambiarlos.")

        def bloque_cliente(col, nombre, clave, pre):
            with col:
                st.markdown(f"**{nombre}**")
                t = st.number_input("Pedidos fecha compromiso totales", min_value=0, step=1, value=int(pre[0]),
                                    key=f"mfc_{clave}_t_{dia_fc}_{origen_pre}")
                e = st.number_input("Entregados", min_value=0, step=1, value=int(pre[1]),
                                    key=f"mfc_{clave}_e_{dia_fc}_{origen_pre}")
                n = st.number_input("No entregados", min_value=0, step=1, value=int(pre[2]),
                                    key=f"mfc_{clave}_n_{dia_fc}_{origen_pre}")
                ok = t == e + n
                if not ok:
                    st.error(f"{nombre}: los totales ({t}) no coinciden con entregados + no entregados ({e + n}).")
                elif t:
                    st.caption(f"NS compromiso {nombre}: {fmt_pct(pct(e, t))}")
            return (t, e, n), ok

        col_e, col_p = st.columns(2)
        val_e, ok_e = bloque_cliente(col_e, "Easy", "e", pre_e)
        val_p, ok_p = bloque_cliente(col_p, "París", "p", pre_p)
        tot_g, ent_g = val_e[0] + val_p[0], val_e[1] + val_p[1]
        if tot_g:
            st.caption(f"NS compromiso general del día: {fmt_pct(pct(ent_g, tot_g))} ({ent_g} de {tot_g} pedidos)")
        b1c, b2c = st.columns(2)
        if b1c.button("💾 Guardar este día", key=f"mfc_g_{dia_fc}", disabled=not (ok_e and ok_p)):
            guardar_fc(dia_fc, val_e, val_p)
            st.success(f"Guardado el {dia_fc:%d/%m/%Y}.")
            st.rerun()
        if dia_fc in ya and ya[dia_fc].get("origen") != "NS diario" and b2c.button(
                "🗑️ Eliminar este día", key=f"mfc_d_{dia_fc}"):
            borrar_fc(dia_fc)
            st.rerun()

        st.markdown("---")
        st.subheader("📈 NS Compromiso del mes")
        dia_c, sem_c, mes_c = tablas_ns_compromiso(regs_m)
        if mes_c is None:
            st.info("Ingresa al menos un día para ver el reporte.")
        else:
            r1, r2, r3 = st.columns(3)
            r1.metric(f"NS Compromiso general de {MESES[mes_m - 1]}", fmt_pct(mes_c["general"]))
            r2.metric("Easy", fmt_pct(mes_c["easy"]) if mes_c["easy"] == mes_c["easy"] else "—")
            r3.metric("París", fmt_pct(mes_c["paris"]) if mes_c["paris"] == mes_c["paris"] else "—")
            st.markdown("**Por semana**")
            mostrar_df(sem_c)
            st.markdown("**Por día**")
            mostrar_df(dia_c)
            st.markdown("**Evolución diaria**")
            graf_c = dia_c.set_index("Fecha")[["NS Easy (%)", "NS París (%)", "NS Compromiso (%)"]]
            try:
                st.line_chart(graf_c, color=["#0068c9", "#83c9ff", "#ff2b2b"])
            except TypeError:  # versiones antiguas de Streamlit sin el parámetro de color
                st.line_chart(graf_c)
            st.caption("Para descargar usa los botones de abajo (Excel o PDF). El ícono pequeño de la tabla baja un "
                       "CSV que Excel en español abre todo en una columna.")
            dc1, dc2 = st.columns(2)
            dc1.download_button("📥 Descargar NS Compromiso en Excel",
                                a_excel(hojas_compromiso(regs_m, anio_m, mes_m), GRAFICO_COMPROMISO_EXCEL),
                                file_name=f"NS_Compromiso_{anio_m}-{mes_m:02d}.xlsx", mime=XLSX, key="dl_comp_xlsx")
            dc2.download_button("📄 Descargar NS Compromiso en PDF", pdf_ns_compromiso(regs_m, anio_m, mes_m),
                                file_name=f"NS_Compromiso_{anio_m}-{mes_m:02d}.pdf", mime="application/pdf",
                                key="dl_comp_pdf")
            dias_btk = set()
            if "acido_dia" in dfs_m and not dfs_m["acido_dia"].empty:
                dias_btk = {datetime.datetime.strptime(x, "%d/%m/%Y").date() for x in dfs_m["acido_dia"]["Fecha"]}
            faltan_fc = sorted(dias_btk - {r["fecha"] for r in regs_m})
            if faltan_fc:
                st.warning("Días con pedidos en el BTK sin datos de compromiso ingresados: "
                           + ", ".join(f"{d:%d/%m}" for d in faltan_fc))
            with st.expander("Ver los datos ingresados"):
                filas_v = []
                for r in sorted(regs_m, key=lambda r: r["fecha"]):
                    fila = {"Fecha": r["fecha"].strftime("%d/%m/%Y"), "Origen": r.get("origen", "Mensual")}
                    for nombre, clave in (("Easy", "easy"), ("París", "paris")):
                        t, e, n = r[clave] if r[clave] is not None else ("—", "—", "—")
                        fila.update({f"{nombre} totales": t, f"{nombre} entregados": e, f"{nombre} no entregados": n})
                    filas_v.append(fila)
                mostrar_df(pd.DataFrame(filas_v))

    # ----- NS Ácido -----
    with sub[2]:
        ac_dia = dfs_m.get("acido_dia", pd.DataFrame())
        if ac_dia.empty:
            st.info(sin_datos)
        else:
            ac_mes = dfs_m["acido_mes"].iloc[0]
            n1, n2, n3, n4 = st.columns(4)
            n1.metric("NS Ácido del mes", fmt_pct(ac_mes["NS Ácido (%)"]))
            n2.metric("Órdenes totales", int(ac_mes["Órdenes totales"]))
            n3.metric("Entregadas / recogidas", int(ac_mes["Entregadas / recogidas"]))
            n4.metric("No entregadas", int(ac_mes["No entregadas"]))
            st.markdown("**Por semana**")
            mostrar_df(dfs_m["acido_semana"])
            st.markdown("**Por día**")
            mostrar_df(ac_dia)
            st.line_chart(ac_dia.set_index("Fecha")[["NS Ácido (%)", "NS Easy (%)", "NS París (%)"]])
            da1, da2 = st.columns(2)
            da1.download_button("📥 Descargar NS Ácido en Excel",
                                a_excel(hojas_acido(dfs_m), GRAFICO_ACIDO_EXCEL),
                                file_name=f"NS_Acido_{anio_m}-{mes_m:02d}.xlsx", mime=XLSX, key="dl_acido_mes")
            da2.download_button("📄 Descargar NS Ácido en PDF", pdf_ns_acido(dfs_m, anio_m, mes_m),
                                file_name=f"NS_Acido_{anio_m}-{mes_m:02d}.pdf", mime="application/pdf",
                                key="dl_acido_pdf")

    # ----- Órdenes por patente -----
    with sub[3]:
        pat_m = dfs_m.get("patentes", pd.DataFrame())
        if pat_m.empty:
            st.info(sin_datos)
        else:
            mostrar_df(pat_m)
            st.caption("% entrega del mes = órdenes entregadas o recogidas / órdenes totales de la patente en el mes "
                       "(solo Easy y París). NS Easy y NS París usan el mismo cálculo con los pedidos de cada cliente.")
            dp1, dp2 = st.columns(2)
            dp1.download_button("📥 Descargar NS por patente en Excel", a_excel({"NS por patente del mes": pat_m}),
                                file_name=f"Ordenes_por_patente_{anio_m}-{mes_m:02d}.xlsx", mime=XLSX, key="dl_pat_mes")
            dp2.download_button("📄 Descargar NS por patente en PDF", pdf_patentes(dfs_m, anio_m, mes_m),
                                file_name=f"NS_por_patente_{anio_m}-{mes_m:02d}.pdf", mime="application/pdf",
                                key="dl_pat_pdf")

    # ----- Descargar todo -----
    with sub[4]:
        st.markdown("Un solo Excel (y un PDF) con todos los reportes del mes: comunas, bonificación, NS compromiso, "
                    "NS ácido, NS por patente y sub-estados de no entrega.")
        if st.button("Preparar reportes del mes", key="b_xls_mes"):
            xls_m = excel_mensual({**dfs_m, **dfs_cal}, regs_m, anio_m, mes_m)
            if xls_m is None:
                st.warning("No hay datos para este mes.")
                st.session_state.pop("excel_mes", None)
                st.session_state.pop("pdf_mes", None)
            else:
                st.session_state["excel_mes"] = (xls_m, f"Reporte_Mensual_{anio_m}-{mes_m:02d}.xlsx")
                st.session_state["pdf_mes"] = (pdf_mensual_completo(dfs_m, regs_m, anio_m, mes_m),
                                               f"Reporte_Mensual_{anio_m}-{mes_m:02d}.pdf")
        if "excel_mes" in st.session_state:
            dt1, dt2 = st.columns(2)
            dt1.download_button("📥 Descargar reporte mensual completo (Excel)", st.session_state["excel_mes"][0],
                                file_name=st.session_state["excel_mes"][1], mime=XLSX, key="dl_mes_todo")
            if "pdf_mes" in st.session_state:
                dt2.download_button("📄 Descargar reporte mensual completo (PDF)", st.session_state["pdf_mes"][0],
                                    file_name=st.session_state["pdf_mes"][1], mime="application/pdf",
                                    key="dl_mes_todo_pdf")

    # ----- Sub-estados de no entregados del mes (al final de la página) -----
    st.markdown("---")
    st.subheader("⚠️ Sub-estados de no entrega del mes")
    if "subestados_noent" not in dfs_m:
        st.info(sin_datos if not dfs_m else "Vuelve a presionar «Procesar mes» para ver los sub-estados de este mes.")
    else:
        se_no = dfs_m["subestados_noent"]
        if se_no.empty:
            st.info("No hubo órdenes no entregadas en el mes.")
        else:
            mostrar_df(se_no)
            st.caption(f"Total de órdenes no entregadas: {int(se_no['Cantidad'].sum())} (solo Easy y París)")
            ds1, ds2 = st.columns(2)
            ds1.download_button("📥 Descargar sub-estados de no entrega en Excel",
                                a_excel({"Sub-estados no entregados": se_no}),
                                file_name=f"Subestados_no_entrega_{anio_m}-{mes_m:02d}.xlsx", mime=XLSX,
                                key="dl_subest_mes")
            ds2.download_button("📄 Descargar sub-estados de no entrega en PDF", pdf_subestados(dfs_m, anio_m, mes_m),
                                file_name=f"Subestados_no_entrega_{anio_m}-{mes_m:02d}.pdf", mime="application/pdf",
                                key="dl_subest_pdf")

    # ----- Calce París del mes -----
    with sub[5]:
        st.markdown("Compara el **BTK Yáñez** con el **BTK París** del mes (ambos con la fecha en la columna A). "
                    "De cada orden se toma su último estado en cada sistema: las que siguen sin coincidir son las "
                    "órdenes que **aún no se arreglan**.")
        if st.button("Procesar calce París del mes", key="b_proc_calce_mes"):
            a_y = resolver_mensual("btk_mensual", primer_m, up_btk_m)
            a_p = resolver_mensual("btk_paris_mensual", primer_m, up_paris_m)
            if a_y is None or a_p is None:
                st.error("Faltan archivos: necesitas el BTK Yáñez del mes y el BTK París del mes.")
            else:
                try:
                    with st.spinner("Cruzando los dos BTK del mes…"):
                        y_mes = preparar_btk(_leer_cache(*a_y))
                        meta_c, dfs_c = procesar_calce_mensual(y_mes, _leer_cache(*a_p), anio_m, mes_m)
                        guardar_resultado("calce_mensual", primer_m, meta_c, dfs_c)
                    st.success(f"¡Calce de {MESES[mes_m - 1]} de {anio_m} procesado!")
                    mostrar_mapeos(meta_c)
                    meta_cal, dfs_cal = meta_c, dfs_c
                except Exception as e:  # noqa: BLE001
                    st.error(f"No se pudo procesar el calce: {e}")
        if not dfs_cal:
            st.info("Todavía no hay calce de este mes. Sube los archivos y presiona «Procesar calce París del mes».")
        else:
            w1, w2, w3, w4 = st.columns(4)
            w1.metric("Órdenes cruzadas", meta_cal["ordenes"])
            w2.metric("Coinciden", meta_cal["coinciden"])
            w3.metric("Aún sin arreglar", meta_cal["pendientes"])
            w4.metric("% coincidencia", fmt_pct(meta_cal["pct_coincidencia"]))
            if meta_cal.get("no_encontradas"):
                st.warning(f"{meta_cal['no_encontradas']} orden(es) de Yáñez no aparecen en el BTK de París.")
            if meta_cal.get("solo_paris"):
                st.caption(f"Hay {meta_cal['solo_paris']} orden(es) en el BTK de París que no están en Yáñez "
                           "(no se cruzan).")
            pend = dfs_cal["calce_pendientes"]
            st.subheader("🚨 Órdenes que aún no se arreglan")
            if pend.empty:
                st.success("✅ Todas las órdenes del mes coinciden.")
            else:
                mostrar_df(estilizar(pend, ["Estado Cruce"], "No Coincide", "Coincide"))
                st.caption("Ordenadas de más a menos días sin arreglarse. La columna «Observación» queda vacía para "
                           "que escribas tus notas en el Excel.")
                c_a, c_b = st.columns(2)
                with c_a:
                    st.markdown("**Por diagnóstico**")
                    mostrar_df(dfs_cal["calce_diagnostico"])
                with c_b:
                    st.markdown("**Por día (en que la orden apareció por primera vez en Yáñez)**")
                    mostrar_df(dfs_cal["calce_dia"])
                with st.expander("Combinaciones de sub-estados que no coinciden (para revisar las reglas)"):
                    mostrar_df(dfs_cal["calce_combinaciones"])
            with st.expander("Ver todas las órdenes cruzadas del mes"):
                mostrar_df(estilizar(dfs_cal["calce_todas"], ["Estado Cruce"], "No Coincide", "Coincide"))
            st.download_button("📥 Descargar calce París del mes en Excel",
                               a_excel({"Pendientes": pend, "Todas las órdenes": dfs_cal["calce_todas"],
                                        "Por día": dfs_cal["calce_dia"], "Diagnóstico": dfs_cal["calce_diagnostico"]}),
                               file_name=f"Calce_Paris_{anio_m}-{mes_m:02d}.xlsx", mime=XLSX, key="dl_calce_mes")
